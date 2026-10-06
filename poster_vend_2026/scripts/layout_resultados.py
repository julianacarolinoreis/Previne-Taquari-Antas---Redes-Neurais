"""Reorganiza Resultados: dispersões empilhadas à esquerda (com PME/EAM/E95 no gráfico),
tempo real no alto da coluna direita e conclusões logo abaixo. Roda depois de revisar_poster.py."""
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

SP = Path(__file__).parent
FIG = SP.parent / "fig"
OUT = SP.parent / "poster_VEND_Santa_Tereza_ALT.pptx"
L, CW, XB = 30.0, 260.0, 304.0


def mm(v):
    return Emu(int(round(v * 36000)))


prs = Presentation(OUT)
slide = prs.slides[0]
S = {sh.name: sh for sh in slide.shapes}


def place(name, x=None, y=None, w=None, h=None):
    sh = S[name]
    if x is not None: sh.left = mm(x)
    if y is not None: sh.top = mm(y)
    if w is not None: sh.width = mm(w)
    if h is not None: sh.height = mm(h)
    return sh


# remove tabela de resultados, sua legenda e a dispersão horizontal
for name in ("Tabela resultados", "Legenda tabela", "Dispersao previsto x observado"):
    el = S[name]._element
    el.getparent().remove(el)

y0 = float(S["Secao RESULTADOS"].top) / 36000
# coluna esquerda: dispersões empilhadas
pic_h = 300 * CW / 260
pic = slide.shapes.add_picture(str(FIG / "dispersao_vertical.png"), mm(L), mm(y0 + 18), width=mm(CW))
pic.name = "Dispersao empilhada"
pic_h = pic.height / 36000
leg = place("Legenda dispersao", x=L, y=y0 + 18 + pic_h + 1, w=CW, h=13)
p = leg.text_frame.paragraphs[0]
p.runs[0].text = ("Figura 2 – Nível previsto × observado (laranja = teste). *Métricas referentes ao conjunto de teste: "
                  "2h – 3 eventos out/2024–2025; 4h – cheia de jun/2025; 8h – cheia de jul/2023.")
for r in p.runs[1:]:
    r._r.getparent().remove(r._r)

# coluna direita: tempo real no topo, gráfico na largura da coluna
place("Secao OPERAÇÃO EM TEMPO REAL", x=XB, y=y0)
fw = CW
fh = fw * 1500 / 3120
place("Operacao ao vivo 2h", x=XB, y=y0 + 18, w=fw, h=fh)
place("Legenda ao vivo", x=XB, y=y0 + 18 + fh + 1, w=CW)
place("Texto ao vivo", x=XB, y=y0 + 18 + fh + 11, w=CW, h=26)

# conclusões abaixo do tempo real
yc = y0 + 18 + fh + 44
place("Secao CONCLUSÕES", x=XB, y=yc)
place("Texto conclusoes", x=XB, y=yc + 19, w=CW, h=66)

# QR no fim da coluna direita, com chamada à esquerda
qs = 56
qy = y0 + 18 + pic_h + 14 - qs - 2
place("Imagem 50", x=XB + CW - qs, y=qy, w=qs, h=qs)
ch = place("Chamada QR code", x=XB + CW - qs - 100, y=qy + qs / 2 - 10, w=94, h=20)
from pptx.enum.text import PP_ALIGN
cp = ch.text_frame.paragraphs[0]
cp.alignment = PP_ALIGN.RIGHT
cp.runs[0].text = "Acompanhe a previsão ao vivo →"
from pptx.util import Pt
cp.runs[0].font.size = Pt(26)
ch.text_frame.word_wrap = True

prs.save(OUT)
print("salvo", OUT, "fim figura", round(y0 + 18 + pic_h + 14))

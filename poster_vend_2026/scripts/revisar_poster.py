"""Aplica a revisão 'de banca' sobre a versão editada pela Juliana (poster_editado_juliana.pptx)."""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

SP = Path(__file__).parent
FIG = SP.parent / "fig"
OUT = SP.parent / "poster_VEND_Santa_Tereza_ALT.pptx"

NAVY = RGBColor(0x2A, 0x37, 0x8D)
ORANGE = RGBColor(0xD9, 0x66, 0x1F)
INK = RGBColor(0x1F, 0x26, 0x33)
MUTED = RGBColor(0x5B, 0x65, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

L, CW, XB = 30.0, 260.0, 304.0


def mm(v):
    return Emu(int(round(v * 36000)))


prs = Presentation(SP / "poster_editado_juliana.pptx")
slide = prs.slides[0]
S = {sh.name: sh for sh in slide.shapes}


def style(run, size, bold=False, color=INK):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = color


def set_runs(para, runs, size, color=INK):
    """Troca o conteúdo do parágrafo preservando pPr (marcadores, espaçamento)."""
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    for t, opt in runs:
        r = para.add_run()
        r.text = t
        style(r, opt.get("size", size), opt.get("bold", False), opt.get("color", color))


def place(name, x=None, y=None, w=None, h=None):
    sh = S[name]
    if x is not None: sh.left = mm(x)
    if y is not None: sh.top = mm(y)
    if w is not None: sh.width = mm(w)
    if h is not None: sh.height = mm(h)
    return sh


# 1) Tabela de configuração: o 2h usa só Santa Tereza e Linha José Júlio
cell = S["Tabela configuracao"].table.cell(1, 3)
r0 = cell.text_frame.paragraphs[0].runs[0]
r0.text = "Santa Tereza e Linha José Júlio"
for r in cell.text_frame.paragraphs[0].runs[1:]:
    r._r.getparent().remove(r._r)

# 2) Introdução: gramática + alinhado à esquerda
intro = place("Texto introducao", x=L, w=CW)
p = intro.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
set_runs(p, [
    ("As cheias de 2023 e 2024 na bacia do Taquari-Antas mostraram a necessidade de previsões de nível "
     "com antecedência útil para o alerta. Este trabalho aplica ", {}),
    ("Redes Neurais Artificiais", {"bold": True}),
    (" do tipo Perceptron Multicamadas (MLP) para prever o nível do rio Taquari em ", {}),
    ("Santa Tereza", {"bold": True}), (" com ", {}), ("2, 4 e 8 horas", {"bold": True}),
    (" de antecedência, na ", {}), ("abordagem alternativa (ALT)", {"bold": True, "color": ORANGE}),
    (": a RNA tem como variável dependente a diferença de nível no horizonte de previsão.", {}),
], 24)
for extra in intro.text_frame.paragraphs[1:]:
    extra._p.getparent().remove(extra._p)

# 3) Legenda do mapa
cap = S["Legenda mapa"].text_frame.paragraphs[0]
set_runs(cap, [("Figura 1 – Bacia do rio Taquari-Antas até Santa Tereza e postos fluviométricos.", {})], 18, MUTED)
place("Legenda mapa", x=325, w=232)

# 4) Títulos de seção alinhados na mesma margem
for name in ("Secao INTRODUÇÃO", "Secao ÁREA DE ESTUDO E DADOS",
             "Secao METODOLOGIA: ABORDAGEM ALTERNATIVA (ALT)", "Secao RESULTADOS", "Secao CONCLUSÕES"):
    place(name, x=L, h=14)
for name in ("Texto area de estudo", "Texto metodologia", "Tabela resultados", "Legenda tabela", "Texto conclusoes",
             "Legenda dispersao"):
    place(name, x=L)

# 5) Figura 2: deixar claro o que é teste
set_runs(S["Legenda dispersao"].text_frame.paragraphs[0], [
    ("Figura 2 – Nível previsto × observado em todas as amostras; em laranja, o conjunto de teste (NSE do título).", {})],
    18, MUTED)

# 6) Tabela 1: quais eventos estão no teste
leg = place("Legenda tabela", y=785, w=CW, h=13)
set_runs(leg.text_frame.paragraphs[0], [
    ("Tabela 1 – Teste (eventos fora do treino): 2h – 3 eventos de out/2024 a 2025 (pico 6,2 m); "
     "4h – cheia de jun/2025 (13,3 m); 8h – cheia de jul/2023 (12,5 m).", {})], 17, MUTED)

# 7) Conclusões com a limitação explícita
place("Secao CONCLUSÕES", y=803)
conc = place("Texto conclusoes", y=822, w=CW, h=70)
ps = conc.text_frame.paragraphs
set_runs(ps[0], [("NSE ≥ 0,979", {"bold": True}), (" e ", {}), ("PME 0,88–0,97", {"bold": True}),
                 (" no teste; erro médio 65–81% menor que a persistência.", {})], 20)
set_runs(ps[1], [("Ao vivo, o 2h acompanhou a cheia de set/2026 (13,3 m) com ", {}),
                 ("EAM de 9 cm", {"bold": True, "color": ORANGE}), (".", {})], 20)
set_runs(ps[2], [("Limitação: ", {"bold": True}),
                 ("as cheias acima de 20 m (2023–2024) ficaram só no treino; o 4h e o 8h seguem em verificação ao vivo.", {})], 20)

# 8) Bloco "tempo real": mesmo padrão de título, figura menor e QR ao lado
sec = place("Secao OPERAÇÃO EM TEMPO REAL", x=XB, y=733, w=196, h=14)
sec.adjustments[0] = 0.5
tf = sec.text_frame
tf.word_wrap = False
tf.margin_left = tf.margin_right = mm(7)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tp = tf.paragraphs[0]
tp.alignment = PP_ALIGN.LEFT
set_runs(tp, [("OPERAÇÃO EM TEMPO REAL", {"bold": True})], 30, WHITE)

fig = place("Operacao ao vivo 2h", x=XB, y=752, w=196, h=196 * 1500 / 3120)
qr_size = 58
qx, qy = XB + CW - qr_size, 760
place("Imagem 50", x=qx, y=qy, w=qr_size, h=qr_size)
chamada = place("Chamada QR code", x=qx - 4, y=qy - 9, w=qr_size + 8, h=8)
cp = chamada.text_frame.paragraphs[0]
cp.alignment = PP_ALIGN.CENTER
set_runs(cp, [("Dashboard ao vivo", {"bold": True})], 18, NAVY)
for extra in chamada.text_frame.paragraphs[1:]:
    extra._p.getparent().remove(extra._p)

legv = place("Legenda ao vivo", x=XB, y=752 + 200 * 1500 / 3120 + 1, w=CW, h=8)
set_runs(legv.text_frame.paragraphs[0],
         [("Figura 3 – Previsão ao vivo 2h × telemetria (25/09 a 05/10/2026).", {})], 18, MUTED)
txt = place("Texto ao vivo", x=XB, y=752 + 200 * 1500 / 3120 + 11, w=CW, h=26)
set_runs(txt.text_frame.paragraphs[0], [
    ("Robô horário com telemetria SGB/ANA, em paralelo ao alerta oficial: ", {}),
    ("230 previsões auditadas", {"bold": True}), (", ", {}),
    ("EAM 9,2 cm", {"bold": True, "color": ORANGE}), (" (persistência 27,2 cm).", {})], 20)

# 9) Rodapé: logos + referências dentro da área branca
ref = place("Referencias e agradecimentos", x=96, y=899, w=468, h=16)
rtf = ref.text_frame
rtf.word_wrap = True
rp = rtf.paragraphs[0]
rp.alignment = PP_ALIGN.LEFT
set_runs(rp, [
    ("Referências: ", {"bold": True}),
    ("ALVISI, S. et al. (2006). Hydrol. Earth Syst. Sci., 10(1), 1-17 · RUMELHART, D. E. et al. (1986). Nature, 323, "
     "533-536 · KUMAR, V. et al. (2023). Sustainability, 15(13), 10543.  ", {}),
    ("Agradecimentos: ", {"bold": True}),
    ("FAPERGS (processos 24/2551-0002124-8 e 25/2551-0002522-2).", {})], 14, MUTED)
for extra in rtf.paragraphs[1:]:
    extra._p.getparent().remove(extra._p)

lh = 17
slide.shapes.add_picture(str(FIG / "ufrgs.png"), mm(L), mm(898), height=mm(lh)).name = "Logo UFRGS"
slide.shapes.add_picture(str(FIG / "fapergs.png"), mm(L + 25), mm(898), height=mm(lh)).name = "Logo FAPERGS"

prs.save(OUT)
print("salvo", OUT)

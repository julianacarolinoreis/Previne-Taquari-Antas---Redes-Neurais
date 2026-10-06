"""Monta o pôster V END / II EBHE sobre o template oficial (594 x 1026 mm)."""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt
from PIL import Image

SP = Path(__file__).parent
FIG = SP.parent / "fig"
OUT = SP.parent / "poster_VEND_Santa_Tereza_ALT.pptx"
M = json.load(open(FIG / "metricas.json"))

NAVY = RGBColor(0x2A, 0x37, 0x8D)
DEEP = RGBColor(0x00, 0x3C, 0x64)
BLUE = RGBColor(0x03, 0x6C, 0xB5)
PALE = RGBColor(0xE6, 0xEE, 0xF5)
ORANGE = RGBColor(0xD9, 0x66, 0x1F)
INK = RGBColor(0x1F, 0x26, 0x33)
MUTED = RGBColor(0x5B, 0x65, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

BODY, SMALL, CAP = 24, 20, 18


def mm(v):
    return Emu(int(round(v * 36000)))


def br(x, nd=3):
    return f"{x:.{nd}f}".replace(".", ",")


prs = Presentation(SP / "template_VEND_594x1026mm.pptx")
slide = prs.slides[0]
shapes = slide.shapes


def _style_run(run, size, bold=False, color=INK, italic=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color


def text(x, y, w, h, paras, size=BODY, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6, line=1.08, name=None):
    """paras: lista de parágrafos; cada um é str, ou lista de (texto, {bold, color, italic, size}),
    ou dict {"runs": ..., "bullet": True}."""
    tb = shapes.add_textbox(mm(x), mm(y), mm(w), mm(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        bullet = False
        if isinstance(p, dict):
            bullet = p.get("bullet", False)
            p = p["runs"]
        if isinstance(p, str):
            p = [(p, {})]
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line
        para.space_after = Pt(space_after)
        if bullet:
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(mm(7)))
            pPr.set("indent", str(-mm(7)))
            bu_clr = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buClr", {})
            srgb = bu_clr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr", {"val": "D9661F"})
            bu_clr.append(srgb)
            bu_font = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buFont", {"typeface": "Arial"})
            bu_char = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "■"})
            pPr.append(bu_clr)
            pPr.append(bu_font)
            pPr.append(bu_char)
        for t, opt in p:
            r = para.add_run()
            r.text = t
            _style_run(r, opt.get("size", size), opt.get("bold", False), opt.get("color", color), opt.get("italic", False))
    return tb


def header(x, y, label, w=None):
    """Rótulo de seção: pílula azul-marinho com texto branco (motivo visual do pôster)."""
    w = w or (len(label) * 7.4 + 22)
    shp = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mm(x), mm(y), mm(w), mm(14))
    shp.name = f"Secao {label}"
    shp.adjustments[0] = 0.5
    shp.fill.solid()
    shp.fill.fore_color.rgb = NAVY
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = mm(7)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = label
    _style_run(r, 30, True, WHITE)
    return shp


def picture(path, x, y, w=None, h=None, name=None):
    iw, ih = Image.open(path).size
    if w and not h:
        h = w * ih / iw
    elif h and not w:
        w = h * iw / ih
    pic = shapes.add_picture(str(path), mm(x), mm(y), mm(w), mm(h))
    if name:
        pic.name = name
    return w, h


def box(x, y, w, h, fill, name, radius=0.12):
    shp = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mm(x), mm(y), mm(w), mm(h))
    shp.name = name
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def table(x, y, w, col_w, rows, size=SMALL, row_h=11.5, name=None, first_col_bold=True):
    nr, nc = len(rows), len(rows[0])
    gt = shapes.add_table(nr, nc, mm(x), mm(y), mm(w), mm(row_h * nr))
    if name:
        gt.name = name
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    # remove o estilo padrão (faixas azuis do Office)
    for k in ("bandRow", "firstRow"):
        tblPr.set(k, "0")
    total = sum(col_w)
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = mm(w * cw / total)
    for i in range(nr):
        tbl.rows[i].height = mm(row_h)
        for j in range(nc):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = mm(2.5)
            c.margin_top = c.margin_bottom = mm(1)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (PALE if i % 2 == 0 else WHITE)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 or (i > 0 and isinstance(rows[i][j], str) and len(rows[i][j]) > 8) else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(rows[i][j])
            _style_run(r, size, bold=(i == 0) or (j == 0 and first_col_bold), color=WHITE if i == 0 else INK)
    return gt


# ------------------------------------------------------------------ cabeçalho
for shp in shapes:
    if shp.name == "TextBox 7":
        shp.left, shp.top, shp.width = mm(30), mm(212), mm(534)
        p = shp.text_frame.paragraphs[0]
        p.line_spacing = None
        pPr = p._p.get_or_add_pPr()
        for ln in pPr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}lnSpc"):
            pPr.remove(ln)
        run = p.runs[0]
        run.text = "Juliana Carolino Reis, Guilherme Garcia de Oliveira, Fernanda Vier Jung e Josiane Leci Vanin Barbieri"
        _style_run(run, 28, False, DEEP)
        for el in ("ea", "cs", "sym"):
            for node in run._r.rPr.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{el}"):
                run._r.rPr.remove(node)
        p2 = shp.text_frame.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(4)
        r2 = p2.add_run()
        r2.text = ("Universidade Federal do Rio Grande do Sul (UFRGS) – Centro Estadual de Pesquisas em "
                   "Sensoriamento Remoto e Meteorologia (CEPSRM)  ·  julianacarolinoreis@gmail.com")
        _style_run(r2, 19, False, MUTED)

L, W = 30.0, 534.0
CW, GAP = 260.0, 14.0
XB = L + CW + GAP
t2, t4, t8 = (M[h]["teste"] for h in ("2h", "4h", "8h"))
av = M["ao_vivo"]["2h"]


def tile(x, y, w, h, big, small, fill=PALE, big_color=NAVY, small_color=INK, big_size=40, name="Destaque"):
    b = box(x, y, w, h, fill, name)
    tf = b.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = mm(3)
    tf.margin_top = tf.margin_bottom = mm(1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = big; _style_run(r, big_size, True, big_color)
    p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = small; _style_run(r, 18, False, small_color)
    return b


# ------------------------------------------------------------------ linha A: introdução + área
y = 252
header(L, y, "INTRODUÇÃO")
text(L, y + 19, CW, 48, [
    [("Previsão do nível do rio Taquari em ", {}), ("Santa Tereza (RS)", {"bold": True}),
     (" com RNA do tipo MLP, ", {}), ("2, 4 e 8 h", {"bold": True}), (" à frente. Na ", {}),
     ("abordagem alternativa (ALT)", {"bold": True, "color": ORANGE}),
     (" a rede prevê só a ", {}), ("variação", {"bold": True}), (" do nível.", {})],
], name="Texto introducao")
ty, tw, th = y + 72, (CW - 2 * 8) / 3, 46
for i, (big, small) in enumerate([("15.800 km²", "bacia até Santa Tereza"),
                                  ("2022–2025", "dados horários ANA/SGB"),
                                  ("3 grupos", "eventos de cheia: treino, validação e teste")]):
    tile(L + i * (tw + 8), ty, tw, th, big, small, big_size=30, name=f"Dado {i + 1}")

mw, mh = picture(FIG / "mapa.png", XB + (CW - 205) / 2, y, w=205, name="Mapa da bacia")
text(XB, y + mh + 2, CW, 8, ["Figura 1 – Bacia do rio Taquari até Santa Tereza e postos fluviométricos."],
     size=CAP, color=MUTED, align=PP_ALIGN.CENTER, name="Legenda mapa")

# ------------------------------------------------------------------ linha B: metodologia (diagrama)
y = 410
header(L, y, "ABORDAGEM ALTERNATIVA (ALT)")
yb, hb = y + 20, 44
etapas = [
    ("Entradas", "níveis, diferenças e acelerações (local e montante) + chuva", PALE, INK),
    ("RNA MLP", "1 camada oculta · sigmoide · retropropagação", PALE, INK),
    ("Saída da RNA", "ΔH = H(t+h) − H(t)", PALE, INK),
    ("Nível previsto", "Ĥ(t+h) = H(t) + ΔH", NAVY, WHITE),
]
bw, aw = 118.0, 20.0
for i, (tit, desc, fill, fc) in enumerate(etapas):
    bx = L + i * (bw + aw)
    b = box(bx, yb, bw, hb, fill, f"Etapa {i + 1}")
    tf = b.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = mm(4)
    tf.margin_top = tf.margin_bottom = mm(2)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = tit
    _style_run(r, 26, True, ORANGE if fc == INK else WHITE)
    p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = desc
    _style_run(r, 19, False, fc)
    if i < len(etapas) - 1:
        a = shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, mm(bx + bw + 3), mm(yb + hb / 2 - 6), mm(aw - 6), mm(12))
        a.name = f"Seta {i + 1}"
        a.fill.solid(); a.fill.fore_color.rgb = ORANGE
        a.line.fill.background(); a.shadow.inherit = False
yp = yb + hb + 7
pw = (W - 2 * 10) / 3
for i, hz in enumerate(("2h", "4h", "8h")):
    tile(L + i * (pw + 10), yp, pw, 14, "", "", fill=WHITE, name=f"Config {hz}")
    text(L + i * (pw + 10), yp + 2, pw, 10, [[(f"ALT {hz}: ", {"bold": True, "color": NAVY}),
         (f"{M[hz]['n_inputs']} entradas · {M[hz]['neuronios']} neurônios", {})]],
         size=21, align=PP_ALIGN.CENTER, name=f"Texto config {hz}")

# ------------------------------------------------------------------ linha C: resultados (dispersão)
y = 507
header(L, y, "RESULTADOS NO TESTE")
dw, dh = picture(FIG / "dispersao.png", L + (W - 455) / 2, y + 17, w=455, name="Dispersao previsto x observado")
text(L, y + 17 + dh + 1, W, 8,
     ["Figura 2 – Nível previsto × observado. Métricas no conjunto de teste (eventos não usados no treino)."],
     size=CAP, color=MUTED, align=PP_ALIGN.CENTER, name="Legenda dispersao")

# ------------------------------------------------------------------ linha D: hidrogramas
y = y + 17 + dh + 13
header(L, y, "CHEIA DE JUL/2023 · 8 h")
ew, eh = picture(FIG / "teste_8h.png", L, y + 18, w=CW, name="Evento de teste 8h")
text(L, y + 18 + eh + 1, CW, 8, ["Figura 3 – A RNA antecipa a subida; a persistência chega 8 h atrasada."],
     size=CAP, color=MUTED, name="Legenda teste 8h")
header(XB, y, "AO VIVO · SET/2026 · 2 h")
vw, vh = picture(FIG / "ao_vivo_2h.png", XB, y + 18, w=CW, name="Operacao ao vivo 2h")
text(XB, y + 18 + vh + 1, CW, 8,
     [f"Figura 4 – {av['n']} previsões auditadas: EAM {br(av['EAM'], 1)} cm (persistência {br(av['EAM_persistencia'], 1)} cm)."],
     size=CAP, color=MUTED, name="Legenda ao vivo")

# ------------------------------------------------------------------ linha E: conclusões + QR code
y = y + 18 + eh + 14
header(L, y, "CONCLUSÕES")
text(L, y + 19, 380, 46, [
    {"bullet": True, "runs": [("NSE ≥ 0,98", {"bold": True}), (" e erro ", {}), ("65–81% menor", {"bold": True}),
                              (" que a persistência nos três horizontes.", {})]},
    {"bullet": True, "runs": [("Em operação, o modelo 2h acompanhou a cheia de set/2026 com ", {}),
                              ("erro médio de 9 cm", {"bold": True, "color": ORANGE}), (".", {})]},
], size=BODY, space_after=6, name="Texto conclusoes")
qs = 58
qx, qy = L + W - qs, y
qr = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mm(qx), mm(qy), mm(qs), mm(qs))
qr.name = "Espaco QR code"
qr.adjustments[0] = 0.06
qr.fill.solid(); qr.fill.fore_color.rgb = WHITE
qr.line.color.rgb = NAVY; qr.line.width = Pt(2.5); qr.line.dash_style = 4
qr.shadow.inherit = False
tf = qr.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "QR CODE"; _style_run(r, 22, True, MUTED)
text(qx - 92, qy + 12, 86, 36, [
    [("Dashboard", {"bold": True, "color": NAVY, "size": 28})],
    [("em tempo real →", {"bold": True, "color": ORANGE, "size": 28})],
], align=PP_ALIGN.RIGHT, name="Chamada QR code")

# ------------------------------------------------------------------ rodapé de referências
text(L, 911, W - 64, 14, [
    [("Referências: ", {"bold": True}),
     ("ALVISI, S. et al. (2006). Hydrology and Earth System Sciences, 10(1), 1-17. · RUMELHART, D. E. et al. (1986). "
      "Nature, 323, 533-536. · KUMAR, V. et al. (2023). Sustainability, 15(13), 10543.   ", {}),
     ("Agradecimentos: ", {"bold": True}),
     ("FAPERGS (processos 24/2551-0002124-8 e 25/2551-0002522-2).", {})],
], size=15, color=MUTED, line=1.0, name="Referencias e agradecimentos")

prs.save(OUT)
print("salvo", OUT)

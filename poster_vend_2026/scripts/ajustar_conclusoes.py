"""Ajusta espaçamento das conclusões na versão v2 da autora, sem mudar o texto."""
from pathlib import Path
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu

SP = Path(__file__).parent
OUT = SP.parent / "poster_VEND_Santa_Tereza_ALT.pptx"
mm = lambda v: Emu(int(round(v * 36000)))
prs = Presentation(SP / "poster_editado_juliana_v2.pptx")
S = {sh.name: sh for sh in prs.slides[0].shapes}
XB, CW = 304.0, 260.0

# texto das conclusões na largura da coluna, alinhado à esquerda (sem os "buracos" do justificado)
c = S["Texto conclusoes"]
sec = S["Secao CONCLUSÕES"]
sec.top = mm(732)
c.left, c.top, c.width, c.height = mm(XB), mm(751), mm(CW), mm(62)
for p in c.text_frame.paragraphs:
    p.alignment = PP_ALIGN.LEFT

# QR abaixo das conclusões, com a chamada à esquerda
qs = 62
qr = S["Imagem 45"]
qr.left, qr.top, qr.width, qr.height = mm(XB + CW - qs), mm(836), mm(qs), mm(qs)
ch = S["Chamada QR code"]
ch.left, ch.top, ch.width, ch.height = mm(XB + CW - qs - 110), mm(836 + qs / 2 - 12), mm(104), mm(24)
for p in ch.text_frame.paragraphs:
    p.alignment = PP_ALIGN.RIGHT

# espaço duplo no texto ao vivo
for p in S["Texto ao vivo"].text_frame.paragraphs:
    prev = None
    for r in p.runs:
        r.text = r.text.replace("  ", " ")
        if prev is not None and prev.text.endswith(" ") and r.text.startswith(" "):
            r.text = r.text.lstrip(" ")
        if r.text:
            prev = r

prs.save(OUT)
print("salvo", OUT)

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Validação de campo do HAND hidráulico LiDAR de Santa Tereza (campanha 05/09/2026).

Lê a planilha de campo (`sumario_erros_hand.xlsx`, aba "versão HAND 2165 CELULAR",
tabela que começa na linha 46) e recalcula, com fórmulas únicas para todos os
pontos e eventos:

1. erro de lâmina (celular): lâmina HAND − lâmina observada, ambas sobre o LiDAR
   na posição registrada pelo celular;
2. erro de cota da superfície d'água com o RTK:
   (z_LiDAR + lâmina HAND) − (z_RTK + lâmina observada);
3. diferença de terreno z_LiDAR(celular) − z_RTK;
4. erros horizontais dos pontos de limite (celular e RTK), como na planilha;
5. validação cruzada leave-one-out do zero vertical (1,60 m) com setembro/2023;
6. conferências com dados do próprio repositório, nas coordenadas RTK:
   ondulação geoidal hgeoHNOR2020, LiDAR ~10 m e HAND web (contornos 0,1 m).

A planilha original não é versionada aqui porque contém coordenadas dos pontos.
Este script grava somente a tabela derivada sem coordenadas, as métricas e as
figuras.

Uso:
  python pesquisas/validacao-campo-hand-santa-tereza/analise_erros_campo.py \
      caminho/para/sumario_erros_hand.xlsx
"""
from __future__ import annotations

import csv
import json
import math
import sys
from pathlib import Path

import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
DATA = ROOT / "assets" / "data"
OUT_CSV = HERE / "erros_campo_santa_tereza.csv"
OUT_JSON = HERE / "metricas_validacao_campo.json"
FIG_DIR = HERE / "figuras"

ZERO_REGUA_M = 1.60          # régua 1,60 m = HAND 0 (versão "2205")
ZERO_REGUA_ANTIGO_M = 2.00   # versão "2165" (lâmina modelo1 da planilha)
GEOIDE_CONST_M = 6.46        # conversão h → H usada na planilha
PICO_REGUA_M = {             # pico da régua 86472600 usado em cada mancha
    "2023-09": 23.65,
    "2023-11": 21.61,
    "2024-05": 22.33,
}
LIMIAR_HAND = {ev: round(v - ZERO_REGUA_M, 2) for ev, v in PICO_REGUA_M.items()}
SHEET = "versão HAND 2165 CELULAR"
HEADER_ROW = 46

# Rótulos RTK reatribuídos na planilha: a cota usada em cada linha corresponde
# ao vértice RTK com outro nome. Mantido explícito para auditoria.
RTK_ALIAS = {"H002": "V009", "V009": "V009A"}


def fnum(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def ler_planilha(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET]
    # vértices RTK (linhas 21–43): nome, N, E, h elipsoidal
    rtk = {}
    for r in range(21, 44):
        nome = ws[f"C{r}"].value
        if not nome:
            continue
        n_mm, e_mm, h_mm = ws[f"D{r}"].value, ws[f"F{r}"].value, ws[f"G{r}"].value
        rtk[str(nome).strip()] = {
            "N": n_mm / 1000.0,
            "E": e_mm / 1000.0,
            "h": h_mm / 1000.0,
            "H": h_mm / 1000.0 - GEOIDE_CONST_M,
        }
    hdr = {}
    for c in ws[HEADER_ROW]:
        if c.value:
            hdr[str(c.value).strip()] = c.column_letter
    linhas = []
    for r in range(HEADER_ROW + 1, HEADER_ROW + 30):
        nome = ws[f"C{r}"].value
        if not nome or not str(nome).strip():
            break

        def g(col_name, _r=r):
            return fnum(ws[f"{hdr[col_name]}{_r}"].value)

        linhas.append({
            "id": str(nome).strip(),
            "altura": g("ALTURA"),
            "lim_ev": (ws[f"{hdr['LIM_EV']}{r}"].value or "").strip(),
            "h_set23": g("H_SET23"),
            "h_nov23": g("H_NOV23"),
            "h_mai24": g("H_MAI24"),
            "z_lidar": g("COTA LIDAR"),
            "z_rtk_planilha": g("COTA RTK SOLO"),
            "lam_set23": g("LAMINA HAND 2205 (-160)"),
            "lam_mai24_planilha": g("LAMINA HAND MAIO/24"),
            "lam_nov23_planilha": g("LAMINA HAND NOV/23"),
            "eh_cel": g("ERRO HORIZONTAL SET 23"),
            "eh_rtk": g("ERRO HORIZONTAL RTK"),
            "rtk_planilha": g("VALIDACAO_RTK_SET23 ERRO VERTICAL"),
            "rtk_planilha_mai": g("ERRO RTK MAI 24"),
            "rtk_planilha_nov": fnum(ws[f"AL{r}"].value),
            "obs": ws[f"Y{r}"].value,
        })
    return rtk, linhas


def montar_registros(rtk, linhas):
    """Uma linha por ponto × evento, com fórmulas únicas."""
    regs, excluidos = [], []
    for p in linhas:
        if p["z_lidar"] is None or p["lam_set23"] is None:
            continue
        if p["obs"]:
            # H010, H011, H012 ("LINHA JOSÉ JULIO") e V012 ("ALTURA MÍNIMA"):
            # fora da tabela de erros da planilha original.
            excluidos.append({"id": p["id"], "motivo": str(p["obs"]).strip()})
            continue
        tipo = "limite" if p["altura"] == 0 else "marca"
        eventos = []
        if tipo == "limite":
            ev = "2024-05" if "mai 2024" in p["lim_ev"] else "2023-09"
            eventos.append((ev, 0.0))
        else:
            for ev, k in (("2023-09", "h_set23"), ("2023-11", "h_nov23"), ("2024-05", "h_mai24")):
                if p[k] is not None and p[k] >= 0:
                    eventos.append((ev, p[k]))
        nome_rtk = RTK_ALIAS.get(p["id"], p["id"])
        v = rtk.get(nome_rtk)
        for ev, lam_obs in eventos:
            lam_hand = p["lam_set23"] - (LIMIAR_HAND["2023-09"] - LIMIAR_HAND[ev])
            # confere com as colunas da planilha (NOV/23 e MAIO/24)
            ref = {"2023-11": p["lam_nov23_planilha"], "2024-05": p["lam_mai24_planilha"]}.get(ev)
            if ref is not None:
                assert abs(ref - lam_hand) < 0.02, (p["id"], ev, ref, lam_hand)
            r = {
                "id": p["id"] + ("b" if p["id"] == "V001" and ev == "2023-11" else ""),
                "tipo": tipo,
                "evento": ev,
                "calibracao": ev == "2023-09",
                "lamina_obs_m": lam_obs,
                "lamina_hand_m": round(lam_hand, 3),
                "erro_lamina_m": round(lam_hand - lam_obs, 3),
                "erro_lamina_zero200_m": round(lam_hand - 0.40 - lam_obs, 3) if ev == "2023-09" else None,
                "z_lidar_m": p["z_lidar"],
                "z_rtk_m": None,
                "dz_lidar_rtk_m": None,
                "erro_cota_rtk_m": None,
                "erro_h_cel_m": p["eh_cel"] if tipo == "limite" else None,
                "erro_h_rtk_m": p["eh_rtk"] if tipo == "limite" else None,
                "rtk_valido": False,
                "nota": "",
            }
            if v is not None:
                z_rtk = v["H"]
                r["z_rtk_m"] = round(z_rtk, 3)
                r["dz_lidar_rtk_m"] = round(p["z_lidar"] - z_rtk, 3)
                r["erro_cota_rtk_m"] = round((p["z_lidar"] + lam_hand) - (z_rtk + lam_obs), 3)
                r["rtk_valido"] = True
                r["_E"], r["_N"], r["_h"] = v["E"], v["N"], v["h"]
            if p["id"] == "V001":
                h1 = rtk["H001"]
                d = math.hypot(v["E"] - h1["E"], v["N"] - h1["N"])
                r["rtk_valido"] = False
                r["nota"] = f"vértice RTK a {d*1000:.0f} mm do H001 (limite); conferir"
            if nome_rtk != p["id"]:
                r["nota"] = f"cota RTK do vértice rotulado {nome_rtk}"
            regs.append(r)
    return regs, excluidos


def metricas(vals):
    a = np.array([x for x in vals if x is not None and not math.isnan(x)], dtype=float)
    if a.size == 0:
        return None
    return {
        "n": int(a.size),
        "vies_m": round(float(a.mean()), 3),
        "mae_m": round(float(np.abs(a).mean()), 3),
        "rmse_m": round(float(np.sqrt((a ** 2).mean())), 3),
        "dp_m": round(float(a.std(ddof=1)), 3) if a.size > 1 else None,
        "mediana_abs_m": round(float(np.median(np.abs(a))), 3),
        "max_abs_m": round(float(np.abs(a).max()), 3),
        "frac_abs_le_0_5m": round(float((np.abs(a) <= 0.5).mean()), 3),
    }


def loocv_zero(regs):
    """Erro fora da amostra se o zero fosse ajustado pela média dos demais pontos."""
    e200 = np.array([r["erro_lamina_zero200_m"] for r in regs if r["calibracao"]])
    n = e200.size
    cv = np.array([e200[i] - (e200.sum() - e200[i]) / (n - 1) for i in range(n)])
    zero_ajustado = ZERO_REGUA_ANTIGO_M + float(e200.mean())
    return cv, zero_ajustado, e200


def conferencias_raster(regs, rtk):
    """Amostra geoide, LiDAR 10 m e HAND web nas coordenadas RTK."""
    out = {}
    try:
        import rasterio
        from pyproj import Transformer
        from scipy.ndimage import map_coordinates
        from PIL import Image
        from shapely.geometry import Point, shape
        from shapely.prepared import prep
    except ImportError as exc:  # pragma: no cover
        print("conferências raster ignoradas:", exc)
        return out
    to_geo = Transformer.from_crs("EPSG:31982", "EPSG:4674", always_xy=True)

    with rasterio.open(DATA / "geoide" / "ondulacao_geoidal_ibge_hnor2020_imbituba.tif") as ds:
        g = ds.read(1).astype(float)
        T = ds.transform
        ns = []
        for nome, v in rtk.items():
            lon, lat = to_geo.transform(v["E"], v["N"])
            col, row = ~T * (lon, lat)
            ns.append(float(map_coordinates(g, [[row - 0.5], [col - 0.5]], order=1)[0]))
        out["geoide_hnor2020_m"] = {"min": round(min(ns), 3), "max": round(max(ns), 3),
                                    "constante_planilha": GEOIDE_CONST_M}

    meta = json.loads((DATA / "santa_tereza_inundacao" / "mdt" / "altitude_terreno_lidar_10m.json").read_text())
    img = np.array(Image.open(DATA / "santa_tereza_inundacao" / "mdt" / meta["png"]))
    z10 = (img[..., 0].astype(float) * 256 + img[..., 1]) * meta["escala"]
    z10[img[..., 3] == 0] = np.nan

    def lidar10(lon, lat):
        c = int((lon - meta["W"]) / (meta["E"] - meta["W"]) * meta["cols"])
        r = int((meta["N"] - lat) / (meta["N"] - meta["S"]) * meta["rows"])
        return float(z10[r, c])

    contornos = json.loads((DATA / "santa_tereza_inundacao" / "contornos_mancha.json").read_text())
    niveis = sorted(((f["properties"]["nivel_m"], prep(shape(f["geometry"])))
                     for f in contornos["features"]), key=lambda t: t[0])

    def hand_web(lon, lat):
        pt = Point(lon, lat)
        lo, hi = 0, len(niveis) - 1
        if not niveis[hi][1].contains(pt):
            return None
        while lo < hi:  # menor nível cujo polígono acumulado contém o ponto
            mid = (lo + hi) // 2
            if niveis[mid][1].contains(pt):
                hi = mid
            else:
                lo = mid + 1
        return niveis[lo][0]

    for r in regs:
        if "_E" not in r:
            continue
        lon, lat = to_geo.transform(r["_E"], r["_N"])
        r["z_lidar10_no_rtk_m"] = round(lidar10(lon, lat), 2)
        hw = hand_web(lon, lat)
        r["hand_web_no_rtk_m"] = hw
        if hw is not None:
            lam = LIMIAR_HAND[r["evento"]] - hw
            r["erro_lamina_hand_web_rtk_m"] = round(lam - r["lamina_obs_m"], 2)
    return out


def figuras(regs, cv):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    FIG_DIR.mkdir(exist_ok=True)
    azul, laranja, verde, cinza = "#2a78d6", "#eb6834", "#1baf7a", "#8a8985"
    plt.rcParams.update({"font.size": 9, "font.family": "DejaVu Sans", "axes.spines.top": False,
                         "axes.spines.right": False, "axes.edgecolor": "#52514e",
                         "axes.labelcolor": "#0b0b0b", "xtick.color": "#52514e", "ytick.color": "#52514e"})

    # Fig. 1 — lâmina observada × modelada (marcas) e erro dos limites
    marcas = [r for r in regs if r["tipo"] == "marca"]
    fig, ax = plt.subplots(figsize=(3.4, 3.2), dpi=300)
    lim = 6.0
    ax.plot([0, lim], [0, lim], color=cinza, lw=1, zorder=1)
    ax.fill_between([0, lim], [-0.5, lim - 0.5], [0.5, lim + 0.5], color=cinza, alpha=0.12, lw=0, zorder=0)
    estilos = {"2023-09": (azul, "o", "set/2023 (calibração)"),
               "2023-11": (laranja, "s", "nov/2023 (independente)"),
               "2024-05": (verde, "^", "mai/2024 (independente)")}
    for ev, (cor, mk, rot) in estilos.items():
        sel = [r for r in marcas if r["evento"] == ev]
        ax.scatter([r["lamina_obs_m"] for r in sel], [r["lamina_hand_m"] for r in sel], s=22, marker=mk,
                   color=cor, edgecolor="white", linewidth=0.6, label=rot, zorder=3)
    for r in marcas:
        if abs(r["erro_lamina_m"]) > 1.0:
            ax.annotate(r["id"], (r["lamina_obs_m"], r["lamina_hand_m"]), xytext=(-26, -3),
                        textcoords="offset points", fontsize=7, color="#52514e")
    ax.set_xlim(0, lim)
    ax.set_ylim(0, lim)
    ax.set_xlabel("Lâmina observada em campo (m)")
    ax.set_ylabel("Lâmina estimada pelo HAND (m)")
    ax.grid(color="#e4e3df", lw=0.5)
    ax.legend(frameon=False, fontsize=7, loc="upper left")
    ax.text(lim - 0.1, lim - 1.35, "faixa ±0,5 m", ha="right", fontsize=7, color="#52514e")
    fig.tight_layout()
    fig.savefig(FIG_DIR / "fig1_lamina_observada_modelada.png")
    plt.close(fig)

    # Fig. 2 — mesmo ponto, duas referências: erro de lâmina (celular+LiDAR) × erro de cota (RTK)
    sel = sorted([r for r in regs if r["rtk_valido"]], key=lambda r: r["erro_lamina_m"])
    fig, ax = plt.subplots(figsize=(3.4, 3.6), dpi=300)
    y = np.arange(len(sel))
    for i, r in enumerate(sel):
        ax.plot([r["erro_lamina_m"], r["erro_cota_rtk_m"]], [i, i], color="#cfcdc7", lw=1.2, zorder=1)
    ax.scatter([r["erro_lamina_m"] for r in sel], y, s=20, color=cinza, zorder=2,
               label="celular + LiDAR (erro de lâmina)")
    ax.scatter([r["erro_cota_rtk_m"] for r in sel], y, s=20, color=azul, zorder=3,
               label="RTK (erro de cota da cheia)")
    ax.axvline(0, color="#52514e", lw=0.8)
    ax.set_yticks(y)
    ax.set_yticklabels([f"{r['id']} ({r['evento'][5:]}/{r['evento'][2:4]})" for r in sel], fontsize=6.5)
    ax.set_xlabel("Erro vertical (m); + = HAND superestima")
    ax.grid(axis="x", color="#e4e3df", lw=0.5)
    ax.legend(frameon=False, fontsize=7, loc="lower left", bbox_to_anchor=(-0.05, 1.0), ncol=1)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "fig2_erro_celular_vs_rtk.png")
    plt.close(fig)

    # Fig. 3 — erro horizontal nos limites: celular × RTK (pareado)
    lims = [r for r in regs if r["tipo"] == "limite" and r["erro_h_cel_m"] is not None]
    lims.sort(key=lambda r: r["erro_h_rtk_m"])
    fig, ax = plt.subplots(figsize=(3.4, 2.6), dpi=300)
    y = np.arange(len(lims))
    for i, r in enumerate(lims):
        ax.plot([r["erro_h_cel_m"], r["erro_h_rtk_m"]], [i, i], color="#cfcdc7", lw=1.2, zorder=1)
    ax.scatter([r["erro_h_cel_m"] for r in lims], y, s=20, color=cinza, zorder=2, label="celular")
    ax.scatter([r["erro_h_rtk_m"] for r in lims], y, s=20, color=azul, zorder=3, label="RTK")
    ax.set_yticks(y)
    ax.set_yticklabels([r["id"] for r in lims], fontsize=7)
    ax.set_xlim(0, None)
    ax.set_xlabel("Distância à borda da mancha modelada (m)")
    ax.grid(axis="x", color="#e4e3df", lw=0.5)
    ax.legend(frameon=False, fontsize=7, loc="lower right")
    fig.tight_layout()
    fig.savefig(FIG_DIR / "fig3_erro_horizontal_limites.png")
    plt.close(fig)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    rtk, linhas = ler_planilha(Path(sys.argv[1]))
    regs, excluidos = montar_registros(rtk, linhas)
    conf = conferencias_raster(regs, rtk)
    cv, zero_aj, e200 = loocv_zero(regs)

    cal = [r for r in regs if r["calibracao"]]
    ind = [r for r in regs if not r["calibracao"]]
    rtkv = [r for r in regs if r["rtk_valido"]]
    m = {
        "limiares_hand_m": LIMIAR_HAND,
        "pico_regua_m": PICO_REGUA_M,
        "zero_regua_m": ZERO_REGUA_M,
        "zero_que_zera_vies_set23_m": round(zero_aj, 3),
        "lamina_celular": {
            "set23_zero200": metricas(list(e200)),
            "set23_zero160": metricas([r["erro_lamina_m"] for r in cal]),
            "set23_zero160_marcas": metricas([r["erro_lamina_m"] for r in cal if r["tipo"] == "marca"]),
            "set23_zero160_limites": metricas([r["erro_lamina_m"] for r in cal if r["tipo"] == "limite"]),
            "set23_loocv": metricas(list(cv)),
            "independente_nov23_mai24": metricas([r["erro_lamina_m"] for r in ind]),
            "todos": metricas([r["erro_lamina_m"] for r in regs]),
            "todos_sem_V010": metricas([r["erro_lamina_m"] for r in regs if r["id"] != "V010"]),
        },
        "cota_rtk": {
            "todos": metricas([r["erro_cota_rtk_m"] for r in rtkv]),
            "marcas": metricas([r["erro_cota_rtk_m"] for r in rtkv if r["tipo"] == "marca"]),
            "limites": metricas([r["erro_cota_rtk_m"] for r in rtkv if r["tipo"] == "limite"]),
            "lamina_celular_mesmos_pontos": metricas([r["erro_lamina_m"] for r in rtkv]),
        },
        "dz_lidar_celular_menos_rtk": {
            "todos": metricas([r["dz_lidar_rtk_m"] for r in rtkv]),
            "marcas": metricas([r["dz_lidar_rtk_m"] for r in rtkv if r["tipo"] == "marca"]),
            "limites": metricas([r["dz_lidar_rtk_m"] for r in rtkv if r["tipo"] == "limite"]),
        },
        "dz_lidar10_no_rtk_menos_rtk": metricas(
            [r["z_lidar10_no_rtk_m"] - r["z_rtk_m"] for r in rtkv if r.get("z_lidar10_no_rtk_m") is not None]),
        "erro_horizontal_limites": {
            "celular": metricas([r["erro_h_cel_m"] for r in regs if r["erro_h_cel_m"] is not None]),
            "rtk": metricas([r["erro_h_rtk_m"] for r in regs if r["erro_h_rtk_m"] is not None]),
        },
        "hand_web_no_rtk_erro_lamina": metricas(
            [r.get("erro_lamina_hand_web_rtk_m") for r in rtkv if r.get("erro_lamina_hand_web_rtk_m") is not None]),
        "conferencias": conf,
        "pontos_excluidos": excluidos,
        "planilha_original_tabela_artigos": {
            "lamina_celular_mae_m": 0.639, "lamina_rtk_mae_m": 1.312,
            "horizontal_celular_media_m": 2.399, "horizontal_rtk_media_m": 4.100,
        },
    }
    for r, e in zip(cal, cv):
        r["erro_lamina_loocv_m"] = round(float(e), 3)

    campos = ["id", "tipo", "evento", "calibracao", "lamina_obs_m", "lamina_hand_m", "erro_lamina_m",
              "erro_lamina_zero200_m", "erro_lamina_loocv_m", "z_lidar_m", "z_rtk_m", "dz_lidar_rtk_m",
              "erro_cota_rtk_m", "erro_h_cel_m", "erro_h_rtk_m", "rtk_valido", "z_lidar10_no_rtk_m",
              "hand_web_no_rtk_m", "erro_lamina_hand_web_rtk_m", "nota"]
    with OUT_CSV.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=campos, extrasaction="ignore")
        w.writeheader()
        for r in regs:
            w.writerow(r)
    OUT_JSON.write_text(json.dumps(m, indent=2, ensure_ascii=False), encoding="utf-8")
    figuras(regs, cv)
    print(json.dumps(m, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

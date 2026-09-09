#!/usr/bin/env python3
"""Audit series availability for frozen STZ/Muçum forcings.

Study gate before calibration — does NOT fill gaps or invent zeros.
"""

from __future__ import annotations

import ast
import csv
import html
import json
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "data" / "estudo_bacia_taquari_antas"
RAIN_CSV = ROOT / "assets" / "data" / "chuvas_horarias.csv"
FORCANTES = OUT / "forcantes_stz_mucum_latest.json"
ROBO = ROOT / "previne" / "robo" / "gerar_previsao_ao_vivo.py"
MUCUM_INPUTS = ROOT / "assets" / "data" / "mucum_modelo_inputs.json"
EXCEL_MUC = ROOT / "assets" / "audit_workbooks" / "12H_ALT__MUC_H12_ALT_STC027_M017.xlsx"
CARREIRO_RAW = ROOT / "assets" / "data" / "hec_hms_carreiro_split" / "raw_ana"


def parse_ts(row: dict) -> str:
    return f"{int(row['ANO']):04d}-{int(row['MES']):02d}-{int(row['DIA']):02d} {int(row['HORA']):02d}:00"


def is_present(val: str | None) -> bool:
    if val is None:
        return False
    s = str(val).strip()
    return s not in ("", "None", "nan", "NaN")


def profile_rain_column(rows: list[dict], col: str) -> dict:
    n = len(rows)
    present_idx = [i for i, r in enumerate(rows) if is_present(r.get(col))]
    present = len(present_idx)
    missing = n - present
    monthly: dict[str, dict] = defaultdict(lambda: {"n": 0, "present": 0})
    for r in rows:
        key = f"{int(r['ANO']):04d}-{int(r['MES']):02d}"
        monthly[key]["n"] += 1
        if is_present(r.get(col)):
            monthly[key]["present"] += 1
    # longest consecutive missing run
    longest = 0
    run = 0
    gap_start = None
    worst = None
    for i, r in enumerate(rows):
        if not is_present(r.get(col)):
            if run == 0:
                gap_start = parse_ts(r)
            run += 1
            if run > longest:
                longest = run
                worst = {"hours": run, "start": gap_start, "end": parse_ts(r)}
        else:
            run = 0
    first = parse_ts(rows[present_idx[0]]) if present_idx else None
    last = parse_ts(rows[present_idx[-1]]) if present_idx else None
    monthly_out = []
    for k in sorted(monthly):
        m = monthly[k]
        monthly_out.append(
            {
                "ym": k,
                "hours": m["n"],
                "present": m["present"],
                "pct": round(100.0 * m["present"] / m["n"], 1) if m["n"] else 0.0,
            }
        )
    pct = round(100.0 * present / n, 1) if n else 0.0
    if present == 0:
        verdict = "vazio"
        gate = "bloquear_como_primaria"
    elif pct >= 90:
        verdict = "ok_operacional"
        gate = "usar"
    elif pct >= 70:
        verdict = "com_buracos"
        gate = "usar_com_mascara_nao_imputar_zero"
    else:
        verdict = "fraca"
        gate = "backup_ou_evitar_como_unica"
    return {
        "column": col,
        "rows": n,
        "present": present,
        "missing": missing,
        "pct_present": pct,
        "first_present": first,
        "last_present": last,
        "longest_gap": worst,
        "monthly": monthly_out,
        "verdict": verdict,
        "gate": gate,
    }


def live_level_codes() -> list[str]:
    text = ROBO.read_text(encoding="utf-8")
    m = re.search(r"ESTACOES_NIVEL\s*=\s*(\[[^\]]*\])", text, re.S)
    if not m:
        return []
    try:
        return [str(x) for x in ast.literal_eval(m.group(1))]
    except Exception:
        return []


def excel_level_presence(codes: list[str]) -> dict:
    out = {"path": str(EXCEL_MUC.relative_to(ROOT)) if EXCEL_MUC.exists() else None, "columns_hit": {}}
    if not EXCEL_MUC.exists():
        out["status"] = "workbook_ausente"
        return out
    try:
        import openpyxl
    except ImportError:
        out["status"] = "openpyxl_ausente"
        return out
    wb = openpyxl.load_workbook(EXCEL_MUC, read_only=True, data_only=True)
    headers = []
    if "DADOS" in wb.sheetnames:
        row1 = next(wb["DADOS"].iter_rows(max_row=1, values_only=True))
        headers = [str(x) for x in row1 if x is not None]
    wb.close()
    for code in codes:
        hits = [h for h in headers if code in h]
        out["columns_hit"][code] = hits
    out["status"] = "ok"
    out["n_headers"] = len(headers)
    return out


def event_telemetry_presence(codes: list[str]) -> dict:
    found = {c: [] for c in codes}
    if not CARREIRO_RAW.exists():
        return {"status": "pasta_ausente", "by_code": found}
    for p in sorted(CARREIRO_RAW.glob("telemetry_*.xml")):
        for c in codes:
            if c in p.name:
                found[c].append(p.name)
    return {"status": "ok", "by_code": found, "path": str(CARREIRO_RAW.relative_to(ROOT))}


def build_payload() -> dict:
    forc = json.loads(FORCANTES.read_text(encoding="utf-8"))
    with RAIN_CSV.open(newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))

    rain_map = {
        "86472000": "chuva_86472000",
        "86472600": "chuva_86472600",
        "2851072": "chuva_02851072",
        "A894": "chuva_inmet_A894",
        "432040401A": "chuva_cemaden_4320404010A",  # naming quirk
    }
    rain_profiles = {code: profile_rain_column(rows, col) for code, col in rain_map.items()}
    # also note legacy empty
    rain_profiles["2851044_legado"] = profile_rain_column(rows, "chuva_02851044")

    level_codes = ["86472000", "86507000", "86125500", "86448000", "86472600", "86510000"]
    live = set(live_level_codes())
    mucum_catalog = set()
    if MUCUM_INPUTS.exists():
        try:
            raw = json.loads(MUCUM_INPUTS.read_text(encoding="utf-8"))
            # flexible shapes
            if isinstance(raw, dict):
                for key in ("stations", "postos", "niveis", "level_stations", "codes"):
                    if key in raw and isinstance(raw[key], list):
                        for item in raw[key]:
                            if isinstance(item, str):
                                mucum_catalog.add(item)
                            elif isinstance(item, dict):
                                for k in ("codigo", "code", "id"):
                                    if k in item:
                                        mucum_catalog.add(str(item[k]))
                # also scan stringified
                blob = json.dumps(raw)
                for c in level_codes:
                    if c in blob:
                        mucum_catalog.add(c)
        except Exception:
            pass
    excel = excel_level_presence(level_codes)
    events = event_telemetry_presence(level_codes)

    level_audit = {}
    for code in level_codes:
        sources = []
        if code in live:
            sources.append("previne_robo_ESTACOES_NIVEL_stz")
        if code in mucum_catalog:
            sources.append("mucum_modelo_inputs")
        if excel.get("columns_hit", {}).get(code):
            sources.append("excel_rna_features")
        if events.get("by_code", {}).get(code):
            sources.append(f"hec_event_telemetry_x{len(events['by_code'][code])}")
        if not sources:
            verdict = "sem_serie_local_encontrada"
            gate = "precisa_baixar_ANA_antes_de_calibrar"
        elif "excel_rna_features" in sources or any(s.startswith("hec_event") for s in sources):
            verdict = "presente_em_rna_ou_eventos"
            gate = "usar_com_auditoria_de_lacunas_por_evento"
        else:
            verdict = "so_live_ou_catalogo"
            gate = "sem_arquivo_historico_longo_garantido_neste_repo"
        role = "alvo_ou_forcante"
        if code == "86472600":
            role = "alvo_STZ / forcante_critica_Mucum"
        elif code == "86510000":
            role = "alvo_Mucum (nao esta no ESTACOES_NIVEL do robo STZ)"
        elif code == "86448000":
            role = "secundario"
        level_audit[code] = {
            "role": role,
            "sources": sources,
            "excel_columns": excel.get("columns_hit", {}).get(code, []),
            "event_files": events.get("by_code", {}).get(code, []),
            "in_live_robo_stz": code in live,
            "in_mucum_catalog": code in mucum_catalog,
            "verdict": verdict,
            "gate": gate,
        }

    # Model readiness
    def rain_ready(codes):
        return {
            c: {
                "verdict": rain_profiles[c]["verdict"],
                "pct": rain_profiles[c]["pct_present"],
                "gate": rain_profiles[c]["gate"],
            }
            for c in codes
        }

    stz_rain = ["86472000", "86472600", "2851072"]
    backups = ["A894", "432040401A"]
    stz_level = ["86472000", "86507000", "86125500"]
    muc_level = ["86472600", "86472000", "86507000", "86125500"]

    blockers = []
    if rain_profiles["A894"]["present"] == 0:
        blockers.append("A894 vazio no CSV publicado — backup só via API INMET; ausencia != zero.")
    if rain_profiles["432040401A"]["pct_present"] < 5:
        blockers.append("CEMADEN quase vazio no CSV (coluna com zero extra no código); backup fraco no arquivo local.")
    if rain_profiles["86472600"]["verdict"] == "com_buracos":
        blockers.append("chuva_86472600 com ~24% buracos e começa tarde (out/2023) — mascarar, não imputar zero.")
    if not level_audit["86125500"]["event_files"]:
        blockers.append("86125500 (Prata) sem telemetria de evento HEC local — só live/Excel.")
    if not level_audit["86448000"]["event_files"]:
        blockers.append("86448000 secundário sem telemetria de evento HEC local.")

    ready = {
        "pode_prototipar_RNA_com_mascaras": True,
        "pode_calibrar_HEC_agora": False,
        "motivo_hec_bloqueado": (
            "Estrutura ok, mas calibração HEC exige séries de chuva/nível reconciliadas por evento "
            "sem imputar zero em buracos; A894/CEMADEN locais fracos; Prata sem raw de evento."
        ),
        "proxima_acao_recomendada": (
            "Baixar/auditar nível histórico ANA para 86125500 (e opcional 86448000); "
            "definir máscara de buracos chuva_86472600; escolher caminho RNA vs .basin HEC."
        ),
    }

    return {
        "schema_version": "estudo_series_forcantes_stz_mucum_v1",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "purpose": "auditoria de disponibilidade de series das forcantes congeladas; nao e calibracao",
        "status": "series_auditadas_v1",
        "discipline_rule": (
            "Buraco != zero. Pane/ausencia nao vira chuva zero. "
            "Esta auditoria libera prototipo com mascara; nao libera NSE/HEC calibrado."
        ),
        "sources": {
            "rain_csv": str(RAIN_CSV.relative_to(ROOT)),
            "rain_rows": len(rows),
            "rain_first": parse_ts(rows[0]) if rows else None,
            "rain_last": parse_ts(rows[-1]) if rows else None,
            "forcantes": str(FORCANTES.relative_to(ROOT)),
            "live_robo": str(ROBO.relative_to(ROOT)),
            "excel_probe": excel,
            "event_telemetry": {"status": events["status"], "path": events.get("path")},
        },
        "rain": rain_profiles,
        "level": level_audit,
        "by_model": {
            "santa_tereza": {
                "target": "86472600",
                "rain_primary": rain_ready(stz_rain),
                "rain_backup": rain_ready(backups),
                "level_primary": {c: level_audit[c] for c in stz_level},
            },
            "mucum": {
                "target": "86510000",
                "rain_primary": rain_ready(stz_rain),
                "rain_backup": rain_ready(backups),
                "level_primary": {c: level_audit[c] for c in muc_level},
                "target_level": level_audit["86510000"],
            },
        },
        "blockers": blockers,
        "readiness": ready,
        "naming_quirks": [
            "Código CEMADEN forçado 432040401A ↔ coluna CSV chuva_cemaden_4320404010A (zero extra).",
            "2851072 ↔ coluna chuva_02851072.",
        ],
        "next_steps": [
            ready["proxima_acao_recomendada"],
            "Não calibrar HEC até fechar máscaras de chuva STZ e série Prata histórica.",
            "Se priorizar RNA: usar forçantes ok_operacional/com_buracos com máscara; excluir A894 vazio do treino.",
        ],
        "artifacts": {
            "json": "series_forcantes_stz_mucum_latest.json",
            "html": "series_forcantes_stz_mucum.html",
        },
        "parent": {
            "forcantes": "forcantes_stz_mucum_latest.json",
            "estrutura": "estrutura_stz_mucum_latest.json",
        },
    }


def write_html(payload: dict) -> None:
    rain_rows = "".join(
        f"<tr><td>{html.escape(code)}</td><td>{html.escape(p['column'])}</td>"
        f"<td>{p['pct_present']}%</td><td>{html.escape(p['verdict'])}</td>"
        f"<td>{html.escape(p['gate'])}</td>"
        f"<td>{html.escape(str(p.get('first_present')))} → {html.escape(str(p.get('last_present')))}</td>"
        f"<td>{html.escape(json.dumps(p.get('longest_gap'), ensure_ascii=False))}</td></tr>"
        for code, p in payload["rain"].items()
        if not code.endswith("_legado")
    )
    level_rows = "".join(
        f"<tr><td>{html.escape(code)}</td><td>{html.escape(p['role'])}</td>"
        f"<td>{html.escape(', '.join(p['sources']) or '—')}</td>"
        f"<td>{html.escape(p['verdict'])}</td><td>{html.escape(p['gate'])}</td></tr>"
        for code, p in payload["level"].items()
    )
    blockers = "".join(f"<li>{html.escape(b)}</li>" for b in payload["blockers"])
    page = f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Séries das forçantes · STZ e Muçum</title>
  <style>
    :root {{ --ink:#1a303f; --muted:#5d7380; --line:#d7e4e8; --ok:#1b7a4a; --warn:#9a5b12; --bad:#a33b35; }}
    body {{ margin:0; color:var(--ink); font:16px/1.55 "Source Sans 3",Segoe UI,sans-serif; background:linear-gradient(165deg,#eef7f8,#fff9f2); }}
    main {{ max-width:1100px; margin:auto; padding:28px 16px 64px; }}
    header, section {{ background:#fff; border:1px solid var(--line); border-radius:18px; padding:20px; margin-bottom:14px; box-shadow:0 10px 26px #1a303f12; }}
    h1 {{ margin:0 0 8px; font:700 clamp(28px,4vw,40px)/1.08 "Fraunces",Georgia,serif; }}
    .eyebrow {{ color:var(--ok); font-size:12px; font-weight:800; letter-spacing:.08em; text-transform:uppercase; }}
    .notice {{ border-left:5px solid var(--warn); background:#fff7e8; color:#6d4810; padding:12px 14px; border-radius:10px; margin:10px 0; }}
    .ok {{ border-left-color:var(--ok); background:#eefaf3; color:#145c38; }}
    .bad {{ border-left-color:var(--bad); background:#fff1ef; color:#7a2d28; }}
    table {{ width:100%; border-collapse:collapse; font-size:13px; }}
    th,td {{ padding:7px; border-bottom:1px solid var(--line); text-align:left; vertical-align:top; }}
    th {{ background:#eef6f7; }}
    .muted {{ color:var(--muted); font-size:13px; }}
    a {{ color:#056999; font-weight:700; }}
  </style>
</head>
<body>
<main>
  <header>
    <div class="eyebrow">Estudo · séries v1</div>
    <h1>Auditoria das forçantes — buraco ≠ zero</h1>
    <p class="muted">{html.escape(payload['generated_at_utc'])} · {html.escape(payload['status'])}</p>
    <div class="notice ok"><strong>Disciplina:</strong> {html.escape(payload['discipline_rule'])}</div>
    <div class="notice {'ok' if payload['readiness']['pode_prototipar_RNA_com_mascaras'] else 'bad'}">
      <strong>RNA com máscara:</strong> {"sim" if payload['readiness']['pode_prototipar_RNA_com_mascaras'] else "não"} ·
      <strong>Calibrar HEC agora:</strong> {"não" if not payload['readiness']['pode_calibrar_HEC_agora'] else "sim"}
    </div>
    <div class="notice bad"><strong>HEC:</strong> {html.escape(payload['readiness']['motivo_hec_bloqueado'])}</div>
  </header>

  <section>
    <h2>Chuva · CSV operacional</h2>
    <p class="muted">{html.escape(payload['sources']['rain_csv'])} · {payload['sources']['rain_rows']} h ·
    {html.escape(str(payload['sources']['rain_first']))} → {html.escape(str(payload['sources']['rain_last']))}</p>
    <table>
      <thead><tr><th>Código</th><th>Coluna</th><th>% presente</th><th>Veredito</th><th>Portão</th><th>1º→último</th><th>Maior buraco</th></tr></thead>
      <tbody>{rain_rows}</tbody>
    </table>
  </section>

  <section>
    <h2>Nível · fontes locais</h2>
    <table>
      <thead><tr><th>Código</th><th>Papel</th><th>Fontes</th><th>Veredito</th><th>Portão</th></tr></thead>
      <tbody>{level_rows}</tbody>
    </table>
  </section>

  <section>
    <h2>Bloqueios / avisos</h2>
    <ul>{blockers}</ul>
    <h2>Próximos</h2>
    <ul>{''.join(f'<li>{html.escape(x)}</li>' for x in payload['next_steps'])}</ul>
    <p>
      <a href="series_forcantes_stz_mucum_latest.json">JSON</a> ·
      <a href="forcantes_stz_mucum.html">forçantes</a> ·
      <a href="estrutura_stz_mucum.html">estrutura</a> ·
      <a href="index.html">estudo-base</a>
    </p>
  </section>
</main>
</body>
</html>
"""
    (OUT / "series_forcantes_stz_mucum.html").write_text(page, encoding="utf-8")


def merge(payload: dict) -> None:
    path = OUT / "estudo_bacia_latest.json"
    if path.exists():
        prev = json.loads(path.read_text(encoding="utf-8"))
        prev["series_forcantes_stz_mucum"] = {
            "status": payload["status"],
            "artifacts": payload["artifacts"],
            "pode_calibrar_hec_agora": payload["readiness"]["pode_calibrar_HEC_agora"],
            "pode_prototipar_rna_com_mascaras": payload["readiness"]["pode_prototipar_RNA_com_mascaras"],
            "updated_at_utc": payload["generated_at_utc"],
        }
        prev["next_study_steps_only"] = payload["next_steps"]
        known = prev.setdefault("known_vs_unknown", {}).setdefault("known", [])
        for item in [
            "Séries de chuva das forçantes auditadas no CSV operacional (A894 vazio; STZ chuva com buracos).",
            "Calibração HEC ainda bloqueada; prototipo RNA com máscara possível.",
        ]:
            if item not in known:
                known.append(item)
        path.write_text(json.dumps(prev, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    for name in ("dois_modelos_stz_mucum_latest.json", "estrutura_stz_mucum_latest.json", "forcantes_stz_mucum_latest.json"):
        p = OUT / name
        if not p.exists():
            continue
        data = json.loads(p.read_text(encoding="utf-8"))
        data["series_artifact"] = payload["artifacts"]
        data["next_steps"] = payload["next_steps"]
        if name.startswith("dois_"):
            data["status"] = "series_auditadas_aguardando_caminho_hec_ou_rna"
        p.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    idx = OUT / "index.html"
    if idx.exists():
        text = idx.read_text(encoding="utf-8")
        if "series_forcantes_stz_mucum.html" not in text:
            text = text.replace(
                '<a href="estrutura_stz_mucum.html">estrutura STZ+Muçum</a></p>',
                '<a href="estrutura_stz_mucum.html">estrutura STZ+Muçum</a> ·\n'
                '       <a href="series_forcantes_stz_mucum.html">séries das forçantes</a></p>',
            )
        if "Séries das forçantes" not in text:
            block = """  <section>
    <h2>Séries das forçantes (v1)</h2>
    <div class="notice" style="border-left-color:#9a5b12;background:#fff7e8;color:#6d4810"><strong>Auditoria:</strong>
    Chuva STZ com buracos; A894 vazio no CSV; HEC ainda bloqueado; RNA com máscara possível.
    Ver <a href="series_forcantes_stz_mucum.html">series_forcantes_stz_mucum.html</a>.</div>
  </section>

"""
            text = text.replace(
                "  <section>\n    <h2>Estrutura candidata (v1)</h2>",
                block + "  <section>\n    <h2>Estrutura candidata (v1)</h2>",
            )
        steps = "".join(f"<li>{html.escape(x)}</li>" for x in payload["next_steps"])
        text = re.sub(
            r"(<h2>Proximos passos de ESTUDO \(sem HEC\)</h2>\s*<ul>)(.*?)(</ul>)",
            r"\1" + steps + r"\3",
            text,
            flags=re.S,
        )
        idx.write_text(text, encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = build_payload()
    (OUT / "series_forcantes_stz_mucum_latest.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    write_html(payload)
    merge(payload)
    summary = {
        "ok": True,
        "status": payload["status"],
        "rain": {k: {"pct": v["pct_present"], "verdict": v["verdict"]} for k, v in payload["rain"].items() if not k.endswith("_legado")},
        "hec_now": payload["readiness"]["pode_calibrar_HEC_agora"],
        "rna_mask_ok": payload["readiness"]["pode_prototipar_RNA_com_mascaras"],
        "blockers": payload["blockers"],
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

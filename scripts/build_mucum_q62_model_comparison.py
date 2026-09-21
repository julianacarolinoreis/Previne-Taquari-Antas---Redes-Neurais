#!/usr/bin/env python3
"""Build a leakage-safe, same-row benchmark for the Muçum Q62 package.

The source RNA candidates are evaluated on the intersection of the five test
event/time keys available in every candidate workbook for each horizon. The
simple baselines and XGBoost are trained only on the canonical workbook's
``Treino`` rows and evaluated on that same common test subset. This is a
research benchmark: it does not select or promote a production model.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import warnings
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from sklearn.ensemble import RandomForestRegressor
from sklearn.linear_model import Ridge
from sklearn.metrics import r2_score
from sklearn.neural_network import MLPRegressor
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from xgboost import XGBRegressor


ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT / "assets" / "data" / "mucum_q62_auditoria.json"
WORKBOOKS = ROOT / "assets" / "audit_workbooks"
OUTPUT_JSON = ROOT / "assets" / "data" / "mucum_q62" / "mucum_q62_model_comparison.json"
OUTPUT_CSV = ROOT / "assets" / "data" / "mucum_q62" / "mucum_q62_model_comparison_metrics.csv"

CANONICAL = {8: "032", 12: "037"}
EVENTS = (31, 33, 34, 35, 37)


def number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        result = float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def timestamp(row: dict[str, Any]) -> str:
    return datetime(
        int(row["ANO"]), int(row["MES"]), int(row["DIA"]), int(row["HORA"])
    ).strftime("%Y-%m-%d %H:%M")


def workbook_for(model_id: str, horizon: int) -> Path:
    matches = sorted(WORKBOOKS.glob(f"{horizon}H_ALT__{model_id}.xlsx"))
    if len(matches) != 1:
        raise RuntimeError(f"esperava uma planilha para {model_id}; encontrei {len(matches)}")
    return matches[0]


def read_workbook(model: dict[str, Any]) -> dict[str, Any]:
    model_id = str(model["model_id"])
    horizon = int(str(model["horizon"]).removesuffix("h"))
    path = workbook_for(model_id, horizon)
    features = [str(value) for value in model.get("input_names", [])]
    forbidden = {"NIVEL_FUTURO", "RNA_FINAL", str(model.get("output_column", ""))}
    if forbidden & set(features):
        raise RuntimeError(f"{model_id}: saída futura apareceu entre as entradas")
    reader = load_workbook(path, read_only=True, data_only=True)
    sheet = reader[reader.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    header = [str(value) for value in next(iterator)]
    indexes = {value: index for index, value in enumerate(header) if value and value != "None"}
    required = set(features) | {"ANO", "MES", "DIA", "HORA", "CONJUNTO", "EVENTO", "NIVEL_FUTURO", "RNA_FINAL", "nivel_86510000"}
    missing = sorted(required - set(indexes))
    if missing:
        raise RuntimeError(f"{model_id}: colunas ausentes {missing}")
    rows: list[dict[str, Any]] = []
    skipped = 0
    for values in iterator:
        row = {key: values[index] for key, index in indexes.items() if index < len(values)}
        event = number(row.get("EVENTO"))
        target = number(row.get("NIVEL_FUTURO"))
        current = number(row.get("nivel_86510000"))
        rna = number(row.get("RNA_FINAL"))
        feature_values = [number(row.get(name)) for name in features]
        if event is None or target is None or current is None or rna is None or any(value is None for value in feature_values):
            skipped += 1
            continue
        base = timestamp(row)
        target_stamp = (datetime.fromisoformat(base) + timedelta(hours=horizon)).strftime("%Y-%m-%d %H:%M")
        rows.append({
            "key": f"{int(event)}|{base}",
            "event": int(event),
            "base_timestamp": base,
            "target_timestamp": target_stamp,
            "partition": str(row.get("CONJUNTO")),
            "observed": target,
            "current": current,
            "rna": rna,
            "features": feature_values,
        })
    return {
        "model_id": model_id,
        "horizon": horizon,
        "path": path,
        "sha256": sha256(path),
        "features": features,
        "rows": rows,
        "skipped": skipped,
    }


def common_keys(workbooks: list[dict[str, Any]], partition: str) -> dict[int, list[str]]:
    by_model = []
    for workbook in workbooks:
        by_model.append({row["key"] for row in workbook["rows"] if row["partition"] == partition and row["event"] in EVENTS})
    common = set.intersection(*by_model)
    result: dict[int, list[str]] = defaultdict(list)
    for key in sorted(common):
        event = int(key.split("|", 1)[0])
        result[event].append(key)
    return dict(result)


def metric_rows(rows: list[dict[str, Any]], model_name: str, horizon: int, method: str, contract_id: str) -> dict[str, Any]:
    if not rows:
        raise RuntimeError(f"sem linhas para calcular {model_name} +{horizon} h")
    obs = np.asarray([row["observed"] for row in rows], dtype=float)
    pred = np.asarray([row["predicted"] for row in rows], dtype=float)
    mae = float(np.mean(np.abs(pred - obs)))
    rmse = float(np.sqrt(np.mean((pred - obs) ** 2)))
    bias = float(np.mean(pred - obs))
    sst = float(np.sum((obs - np.mean(obs)) ** 2))
    nse = float(1 - np.sum((pred - obs) ** 2) / sst) if sst else None
    r2 = float(r2_score(obs, pred)) if len(rows) > 1 else None
    observed_peak = max(rows, key=lambda row: row["observed"])
    predicted_peak = max(rows, key=lambda row: row["predicted"])
    peak_error = float(predicted_peak["predicted"] - observed_peak["observed"])
    lag = (datetime.fromisoformat(predicted_peak["target_timestamp"]) - datetime.fromisoformat(observed_peak["target_timestamp"])).total_seconds() / 3600
    return {
        "model": model_name,
        "method": method,
        "horizon_hours": horizon,
        "contract_id": contract_id,
        "event": "all",
        "n": len(rows),
        "mae_cm": round(mae, 6),
        "rmse_cm": round(rmse, 6),
        "bias_cm": round(bias, 6),
        "nse": round(nse, 8) if nse is not None else None,
        "r2": round(r2, 8) if r2 is not None else None,
        "peak_observed_cm": round(float(observed_peak["observed"]), 6),
        "peak_predicted_cm": round(float(predicted_peak["predicted"]), 6),
        "peak_error_cm": round(peak_error, 6),
        "peak_abs_error_cm": round(abs(peak_error), 6),
        "peak_lag_hours": round(float(lag), 6),
        "base_start": rows[0]["base_timestamp"],
        "base_end": rows[-1]["base_timestamp"],
    }


def fit_baselines(train: list[dict[str, Any]], features: int) -> dict[str, Any]:
    x = np.asarray([row["features"] for row in train], dtype=float).reshape(-1, features)
    y = np.asarray([row["observed"] for row in train], dtype=float)
    return {
        "Ridge": make_pipeline(StandardScaler(), Ridge(alpha=10.0)),
        "MLP": make_pipeline(StandardScaler(), MLPRegressor(hidden_layer_sizes=(32, 16), activation="relu", solver="lbfgs", alpha=0.001, max_iter=800, random_state=42)),
        "Random Forest": RandomForestRegressor(n_estimators=220, max_depth=12, min_samples_leaf=2, random_state=42, n_jobs=2),
        "XGBoost": XGBRegressor(n_estimators=240, max_depth=3, learning_rate=0.04, min_child_weight=2, subsample=0.9, colsample_bytree=0.9, objective="reg:squarederror", eval_metric="rmse", random_state=42, n_jobs=2, tree_method="hist"),
    }, x, y


def build() -> dict[str, Any]:
    audit = json.loads(AUDIT.read_text(encoding="utf-8"))
    models = audit["models"]
    by_horizon: dict[int, list[dict[str, Any]]] = defaultdict(list)
    source: dict[str, dict[str, Any]] = {}
    for model in models:
        item = read_workbook(model)
        source[item["model_id"]] = item
        by_horizon[item["horizon"]].append(item)

    event_metrics: list[dict[str, Any]] = []
    aggregate: list[dict[str, Any]] = []
    fairness: dict[str, Any] = {}
    fit_audit: dict[str, dict[str, Any]] = {}
    csv_rows: list[dict[str, Any]] = []

    for horizon in sorted(by_horizon):
        workbooks = sorted(by_horizon[horizon], key=lambda item: item["model_id"])
        keys_by_event = common_keys(workbooks, "Teste")
        canonical_id = next(item["model_id"] for item in workbooks if item["model_id"].startswith(CANONICAL[horizon]))
        canonical = source[canonical_id]
        canonical_by_key = {row["key"]: row for row in canonical["rows"]}
        test_keys = [key for event in EVENTS for key in keys_by_event.get(event, [])]
        train = [row for row in canonical["rows"] if row["partition"] == "Treino"]
        if not train or not test_keys:
            raise RuntimeError(f"contrato vazio no horizonte {horizon}")
        fairness[str(horizon)] = {
            "canonical_model": canonical_id,
            "canonical_workbook": rel(canonical["path"]),
            "canonical_workbook_sha256": canonical["sha256"],
            "feature_count": len(canonical["features"]),
            "feature_names": canonical["features"],
            "train_partition": "Treino",
            "test_partition": "Teste",
            "common_test_points_total": len(test_keys),
            "common_test_points_by_event": {str(event): len(keys_by_event.get(event, [])) for event in EVENTS},
        }

        model_predictions: dict[str, dict[str, float]] = {}
        model_meta: dict[str, tuple[str, str]] = {}
        for workbook in workbooks:
            rows_by_key = {row["key"]: row for row in workbook["rows"]}
            model_name = f"RNA {workbook['model_id'][:3]}"
            model_predictions[model_name] = {key: rows_by_key[key]["rna"] for key in test_keys}
            model_meta[model_name] = ("RNA fonte Q62", workbook["model_id"])

        baseline_models, train_x, train_y = fit_baselines(train, len(canonical["features"]))
        test_rows = [canonical_by_key[key] for key in test_keys]
        model_predictions["Persistência"] = {key: canonical_by_key[key]["current"] for key in test_keys}
        model_meta["Persistência"] = ("baseline", "nivel_86510000")
        test_x = np.asarray([row["features"] for row in test_rows], dtype=float)
        for model_name, estimator in baseline_models.items():
            with warnings.catch_warnings(record=True) as fit_warnings:
                warnings.simplefilter("always")
                estimator.fit(train_x, train_y)
            final_estimator = estimator[-1] if hasattr(estimator, "steps") else estimator
            iterations = getattr(final_estimator, "n_iter_", None)
            loss = getattr(final_estimator, "loss_", None)
            fit_audit[f"{horizon}h::{model_name}"] = {
                "model": model_name,
                "horizon_hours": horizon,
                "train_points": len(train),
                "converged_or_completed": not bool(fit_warnings),
                "warning_count": len(fit_warnings),
                "warnings": [str(item.message) for item in fit_warnings],
                "iterations": int(iterations) if iterations is not None else None,
                "loss": round(float(loss), 8) if loss is not None else None,
            }
            pred = estimator.predict(test_x)
            model_predictions[model_name] = {key: float(value) for key, value in zip(test_keys, pred)}
            model_meta[model_name] = ("baseline", canonical_id)

        model_event_metrics: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for model_name, predictions in model_predictions.items():
            for event in EVENTS:
                rows = []
                for key in keys_by_event.get(event, []):
                    base = canonical_by_key[key]
                    rows.append({**base, "predicted": predictions[key]})
                metrics = metric_rows(rows, model_name, horizon, model_meta[model_name][0], model_meta[model_name][1])
                metrics["event"] = event
                metrics["n_common_event_points"] = len(rows)
                event_metrics.append(metrics)
                model_event_metrics[model_name].append(metrics)
                csv_rows.append(metrics)
            all_rows = []
            for key in test_keys:
                base = canonical_by_key[key]
                all_rows.append({**base, "predicted": predictions[key]})
            pooled = metric_rows(all_rows, model_name, horizon, model_meta[model_name][0], model_meta[model_name][1])
            event_maes = [item["mae_cm"] for item in model_event_metrics[model_name]]
            event_peaks = [item["peak_abs_error_cm"] for item in model_event_metrics[model_name]]
            event_lags = [abs(item["peak_lag_hours"]) for item in model_event_metrics[model_name]]
            aggregate.append({
                **pooled,
                "event_mae_mean_cm": round(float(np.mean(event_maes)), 6),
                "event_mae_median_cm": round(float(np.median(event_maes)), 6),
                "event_mae_worst_cm": round(float(np.max(event_maes)), 6),
                "event_peak_abs_error_mean_cm": round(float(np.mean(event_peaks)), 6),
                "event_peak_abs_error_worst_cm": round(float(np.max(event_peaks)), 6),
                "event_peak_lag_abs_mean_hours": round(float(np.mean(event_lags)), 6),
                "train_points": len(train),
                "common_test_points": len(test_keys),
            })

    for horizon in sorted(by_horizon):
        rows = [item for item in aggregate if item["horizon_hours"] == horizon]
        for metric in ("event_mae_mean_cm", "rmse_cm", "event_peak_abs_error_mean_cm"):
            ranked = sorted(rows, key=lambda item: item[metric])
            for rank, item in enumerate(ranked, start=1):
                item[f"rank_{metric}"] = rank
        event_rows = [item for item in event_metrics if item["horizon_hours"] == horizon]
        by_event = defaultdict(list)
        for item in event_rows:
            by_event[item["event"]].append(item)
        for item in rows:
            item["best_event_mae_count"] = sum(
                min(group, key=lambda candidate: candidate["mae_cm"])["model"] == item["model"]
                for group in by_event.values()
            )

    skipped = [
        {"model": "LSTM/GRU", "status": "not_executed", "reason": "O repositório não contém um runtime temporal auditado para esta partição; instalar uma biblioteca não cria um contrato de sequência nem uma comparação reproduzível."},
        {"model": "Transformer temporal", "status": "not_executed", "reason": "Requer sequência temporal, ajuste de janela e validação por eventos; não deve ser anunciado como comparação justa antes desse protocolo."},
        {"model": "Transformer-GNN", "status": "not_executed", "reason": "Não há ainda um grafo hidrográfico validado e um conjunto multie estação com partição compatível nesta rodada Q62."},
    ]
    output = {
        "schema_version": 1,
        "artifact_id": "mucum_q62_model_comparison",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
        "status": "research_only_same_common_test_rows",
        "research_only": True,
        "official_alert": False,
        "promotion_allowed": False,
        "purpose": "Comparação justa de RNA fonte, persistência, Ridge, MLP, Random Forest e XGBoost usando o mesmo recorte comum de cinco eventos Q62.",
        "method": {
            "train_partition": "Treino",
            "test_partition": "Teste",
            "test_event_count": len(EVENTS),
            "test_events": list(EVENTS),
            "same_common_test_rows": True,
            "baseline_training": "somente Treino da planilha canônica de cada horizonte",
            "target": "NIVEL_FUTURO",
            "units": "cm",
            "peak_lag": "timestamp do pico previsto menos timestamp do pico observado",
            "no_future_features": True,
        },
        "fairness_by_horizon": fairness,
        "aggregate_metrics": sorted(aggregate, key=lambda item: (item["horizon_hours"], item["event_mae_mean_cm"])),
        "event_metrics": event_metrics,
        "fit_audit": fit_audit,
        "skipped_models": skipped,
        "uncertainty": {
            "status": "not_calibrated_for_live_use",
            "research_only": True,
            "available": ["resíduos e erro do pico no teste comum", "métricas por evento e horizonte", "separação temporal Treino/Teste"],
            "not_available": ["intervalos de predição calibrados", "incerteza meteorológica", "incerteza MDT/HAND e nível→mancha", "incerteza de estação/vazão"],
            "interpretation": "O erro observado no replay não é uma faixa de segurança nem uma probabilidade operacional.",
        },
        "audit": {
            "source_workbook_count": len(source),
            "source_workbook_sha256": {item["model_id"]: item["sha256"] for item in source.values()},
            "finite_rows_only": True,
            "duplicate_common_keys": 0,
            "target_and_output_columns_in_inputs": False,
            "source_rna_not_retrained": True,
        },
        "sources": [
            "assets/data/mucum_q62_auditoria.json",
            "assets/data/mucum_q62/mucum_q62_replay_candidates.json",
            "assets/audit_workbooks/",
        ],
    }
    OUTPUT_JSON.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with OUTPUT_CSV.open("w", newline="", encoding="utf-8") as handle:
        fields = sorted({key for row in csv_rows for key in row})
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(csv_rows)
    return output


if __name__ == "__main__":
    result = build()
    print(json.dumps({"output": rel(OUTPUT_JSON), "aggregate_rows": len(result["aggregate_metrics"]), "event_rows": len(result["event_metrics"])}, ensure_ascii=False))

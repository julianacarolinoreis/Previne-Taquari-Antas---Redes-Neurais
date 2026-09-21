#!/usr/bin/env python3
"""Build coherent same-row comparisons for Santa Tereza +4 h and +8 h.

The source catalog contains many rotations, but they do not all share the
same independent event.  This builder therefore publishes two explicit
contracts instead of pretending that all files are one experiment:

* +4 h: the 11 V01--V11 ALT variants with the same R10 split, evaluated on
  the intersection of event 13 / Teste rows;
* +8 h: a selected, documented ALT/CONV family set evaluated on the
  intersection of event 3 / Teste rows.

Source RNA outputs are not retrained.  Ridge, MLP, Random Forest and XGBoost
are trained only on the canonical workbook's Treino rows and tested on the
same common rows.  The artifact is research-only and is never a promotion
decision.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import warnings
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

from build_mucum_q62_model_comparison import fit_baselines, metric_rows


ROOT = Path(__file__).resolve().parents[1]
WORKBOOKS = ROOT / "assets" / "audit_workbooks"
OUTPUT_DIR = ROOT / "assets" / "data" / "santa_tereza_horizons"
OUTPUT_JSON = OUTPUT_DIR / "stz_horizon_model_comparisons.json"
OUTPUT_CSV = OUTPUT_DIR / "stz_horizon_model_comparisons_metrics.csv"

STZ_4H_FILES = sorted(WORKBOOKS.glob("4H_ALT__V*_R10_T19-21_V1-3-5-15-17_*.xlsx"))
STZ_8H_FILES = [
    WORKBOOKS / "8H_ALT__01_8h_alt_8H_ALT_C0169__8a1d4f6c2b.xlsx",
    WORKBOOKS / "8H_ALT__03_8h_alt_8H_ALT_C0265__4f43f568ce.xlsx",
    WORKBOOKS / "8H_ALT__04_8h_alt_8H_ALT_C0273__125b8d65c6.xlsx",
    WORKBOOKS / "8H_ALT__05_8H_ALT_linha003_altR_004_06_8h_alt_8H_ALT_C0217__f48ca4eb07.xlsx",
    WORKBOOKS / "8H_ALT__07_8h_alt_8H_ALT_C0241__dd12e4f12d.xlsx",
    WORKBOOKS / "8H_ALT__08_8h_alt_8H_ALT_C0289__0e65a5a5ac.xlsx",
    WORKBOOKS / "8H_ALT__10_8h_alt_8H_ALT_C0569__7c299ec1cf.xlsx",
    WORKBOOKS / "8H_ALT__11_8h_alt_8H_ALT_C0174__14e7366728.xlsx",
    WORKBOOKS / "8H_CONV__24_8h_conv_8H_CONV_C0289__e4b81346c6.xlsx",
    WORKBOOKS / "8H_CONV__28_8h_conv_8H_CONV_C0078__3e700886ed.xlsx",
]


def number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        result = float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def base_timestamp(values: tuple[Any, ...] | list[Any], indexes: dict[str, int]) -> str:
    year = int(number(values[indexes["ANO"]]) or 0)
    month = int(number(values[indexes["MES"]]) or 0)
    day = int(number(values[indexes["DIA"]]) or 0)
    hour = int(number(values[indexes["HORA"]]) or 0)
    minute = int(number(values[indexes["MINUTO"]]) or 0)
    return datetime(year, month, day, hour, minute).strftime("%Y-%m-%d %H:%M")


def short_id(path: Path, horizon: int) -> str:
    stem = path.name.removesuffix(".xlsx")
    prefix = f"{horizon}H_ALT__" if horizon == 4 else ("8H_ALT__" if "8H_ALT__" in stem else "8H_CONV__")
    return stem.removeprefix(prefix)


def family(path: Path) -> str:
    return "CONV" if "8H_CONV__" in path.name else "ALT"


def read_workbook(path: Path, horizon: int) -> dict[str, Any]:
    if not path.exists():
        raise RuntimeError(f"planilha ausente: {path.name}")
    reader = load_workbook(path, read_only=True, data_only=True)
    sheet_name = reader.sheetnames[0]
    sheet = reader[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    header = list(next(iterator))
    indexes = {str(value): index for index, value in enumerate(header) if value not in (None, "")}
    input_prefix = "input_" if horizon == 4 else "inp"
    feature_names = [str(value) for value in header if str(value).startswith(input_prefix)]
    target_name = "OBSERVADO_CM" if horizon == 4 else "OBSERVADO_CM_AUDITORIA"
    output_name = "RNA_CM"
    required = {"EVENTO", "CONJUNTO", "ANO", "MES", "DIA", "HORA", "MINUTO", target_name, output_name, "NIVEL_ATUAL_CM"}
    missing = sorted(required - set(indexes))
    if missing:
        raise RuntimeError(f"{path.name}: colunas ausentes {missing}")
    rows: list[dict[str, Any]] = []
    skipped = 0
    for values in iterator:
        event = number(values[indexes["EVENTO"]])
        target = number(values[indexes[target_name]])
        current = number(values[indexes["NIVEL_ATUAL_CM"]])
        rna = number(values[indexes[output_name]])
        if event is None or target is None or current is None or rna is None:
            skipped += 1
            continue
        base = base_timestamp(values, indexes)
        features = [number(values[indexes[name]]) for name in feature_names]
        rows.append({
            "key": f"{int(event)}|{base}",
            "event": int(event),
            "base_timestamp": base,
            "target_timestamp": (datetime.fromisoformat(base) + timedelta(hours=horizon)).strftime("%Y-%m-%d %H:%M"),
            "partition": str(values[indexes["CONJUNTO"]]),
            "observed": target,
            "current": current,
            "rna": rna,
            "features": features,
            "features_complete": all(value is not None for value in features),
        })
    return {
        "path": path,
        "sha256": sha256(path),
        "model_id": short_id(path, horizon),
        "family": family(path),
        "feature_names": feature_names,
        "rows": rows,
        "skipped": skipped,
        "sheet": sheet_name,
    }


def model_label(workbook: dict[str, Any], horizon: int) -> str:
    model_id = workbook["model_id"]
    if horizon == 4:
        match = re.search(r"V\d+", model_id)
        return f"RNA STZ 4H {match.group(0) if match else model_id[:5]}"
    match = re.search(r"C\d+", model_id)
    return f"RNA STZ 8H {workbook['family']} {match.group(0) if match else model_id[:12]}"


def build_contract(horizon: int, paths: list[Path], event: int, canonical_name_part: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    source = [read_workbook(path, horizon) for path in paths]
    common_sets = []
    for workbook in source:
        common_sets.append({row["key"] for row in workbook["rows"] if row["partition"] == "Teste" and row["event"] == event})
    common = set.intersection(*common_sets)
    if not common:
        raise RuntimeError(f"+{horizon} h: não há interseção do evento {event} / Teste")
    test_keys = sorted(common)
    canonical = next((item for item in source if canonical_name_part in item["model_id"]), None)
    if canonical is None:
        raise RuntimeError(f"+{horizon} h: canonical não encontrado: {canonical_name_part}")
    canonical_by_key = {row["key"]: row for row in canonical["rows"]}
    train = [row for row in canonical["rows"] if row["partition"] == "Treino" and row["features_complete"]]
    if not train:
        raise RuntimeError(f"+{horizon} h: canonical sem linhas Treino completas")
    test_rows = [canonical_by_key[key] for key in test_keys]
    model_predictions: dict[str, dict[str, float]] = {}
    model_meta: dict[str, tuple[str, str]] = {}
    for workbook in source:
        by_key = {row["key"]: row for row in workbook["rows"]}
        label = model_label(workbook, horizon)
        model_predictions[label] = {key: by_key[key]["rna"] for key in test_keys}
        model_meta[label] = (f"RNA fonte STZ {horizon}H {workbook['family']}", workbook["model_id"])

    model_predictions["Persistência"] = {key: canonical_by_key[key]["current"] for key in test_keys}
    model_meta["Persistência"] = ("baseline", "NIVEL_ATUAL_CM")
    train_x = np.asarray([row["features"] for row in train], dtype=float)
    train_y = np.asarray([row["observed"] for row in train], dtype=float)
    test_x = np.asarray([row["features"] for row in test_rows], dtype=float)
    baseline_models, _, _ = fit_baselines(train, len(canonical["feature_names"]))
    fit_audit: dict[str, Any] = {}
    for name, estimator in baseline_models.items():
        with warnings.catch_warnings(record=True) as fit_warnings:
            warnings.simplefilter("always")
            estimator.fit(train_x, train_y)
        final_estimator = estimator[-1] if hasattr(estimator, "steps") else estimator
        fit_audit[name] = {
            "train_points": len(train),
            "feature_count": len(canonical["feature_names"]),
            "warnings": [str(item.message) for item in fit_warnings],
            "warning_count": len(fit_warnings),
            "iterations": getattr(final_estimator, "n_iter_", None),
        }
        model_predictions[name] = {key: float(value) for key, value in zip(test_keys, estimator.predict(test_x))}
        model_meta[name] = ("baseline", canonical["model_id"])

    metrics = []
    for name, predictions in model_predictions.items():
        rows = [{**canonical_by_key[key], "predicted": predictions[key]} for key in test_keys]
        rows.sort(key=lambda row: row["base_timestamp"])
        item = metric_rows(rows, name, horizon, model_meta[name][0], model_meta[name][1])
        item.update({"event": event, "independent_partition": "Teste", "n_common_event_points": len(rows), "family": "baseline" if name in {"Persistência", "Ridge", "MLP", "Random Forest", "XGBoost"} else model_meta[name][0].rsplit(" ", 1)[-1]})
        metrics.append(item)
    metrics.sort(key=lambda item: item["mae_cm"])
    excluded = []
    for workbook, keys in zip(source, common_sets):
        if len(keys) != len(common):
            excluded.append({"model": workbook["model_id"], "status": "partial_common_event", "event_points": len(keys), "common_points": len(common)})
    contract = {
        "horizon_hours": horizon,
        "event": event,
        "partition": "Teste",
        "same_common_test_rows": True,
        "common_points": len(test_keys),
        "canonical_model": canonical["model_id"],
        "canonical_workbook": rel(canonical["path"]),
        "canonical_workbook_sha256": canonical["sha256"],
        "train_partition": "Treino",
        "feature_count": len(canonical["feature_names"]),
        "feature_names": canonical["feature_names"],
        "target": "observed level in cm",
        "no_future_features": True,
        "source_models": [
            {"model_id": workbook["model_id"], "family": workbook["family"], "workbook": rel(workbook["path"]), "sha256": workbook["sha256"], "label": model_label(workbook, horizon)}
            for workbook in source
        ],
        "excluded_or_partial": excluded,
        "aggregate_metrics": metrics,
        "fit_audit": fit_audit,
        "audit": {
            "source_rna_not_retrained": True,
            "common_key_intersection": True,
            "train_and_test_temporally_separated": True,
            "target_and_output_columns_not_used_as_features": True,
            "flow_calibration_status": "blocked_for_flow_calibration",
            "flow_calibration_reason": "Santa Tereza mantém nível em cm; não existe neste contrato vazão horária reconciliada com curva-chave oficial.",
        },
    }
    return contract, metrics


def build() -> dict[str, Any]:
    four, four_csv = build_contract(4, STZ_4H_FILES, 13, "V11_")
    eight, eight_csv = build_contract(8, STZ_8H_FILES, 3, "08_8h_alt")
    output = {
        "schema_version": 1,
        "artifact_id": "santa_tereza_horizon_model_comparisons",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
        "status": "research_only_same_common_independent_event_by_horizon",
        "research_only": True,
        "official_alert": False,
        "promotion_allowed": False,
        "station": {"code": "86472600", "name": "SANTA TEREZA", "river": "RIO TAQUARI"},
        "scope": {
            "four_hours": "11 variantes 4H_ALT V01-V11 do mesmo recorte R10; evento 13 / Teste; interseção publicada.",
            "eight_hours": "8 variantes 8H_ALT e 2 variantes 8H_CONV selecionadas por contrato comum; evento 3 / Teste; não representa as 190 entradas do catálogo inteiro.",
            "two_hours": "ver assets/data/santa_tereza_2h/stz_2h_model_comparison.json; E12 / Verificacao.",
        },
        "horizons": {"4": four, "8": eight},
        "uncertainty": {
            "status": "not_calibrated_for_live_use",
            "research_only": True,
            "available": [
                "resíduos pontuais e erro do pico no recorte independente",
                "comparação entre fonte RNA, persistência e baselines",
                "separação explícita entre Treino e Teste",
            ],
            "not_available": [
                "intervalos de predição calibrados",
                "incerteza meteorológica probabilística",
                "incerteza MDT/HAND e propagação nível→mancha",
                "conversão validada de nível em vazão para Santa Tereza",
            ],
            "interpretation": "Os números são diagnóstico de replay; não são faixa de segurança, alerta ou probabilidade operacional.",
        },
        "sources": ["assets/data/auditaveis_series.json", "assets/audit_workbooks/", "assets/data/santa_tereza_2h/stz_2h_model_comparison.json"],
    }
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_JSON.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    rows = []
    for horizon_rows in (four_csv, eight_csv):
        rows.extend(horizon_rows)
    with OUTPUT_CSV.open("w", newline="", encoding="utf-8") as handle:
        fields = sorted({key for row in rows for key in row})
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    return output


if __name__ == "__main__":
    data = build()
    print(json.dumps({"output": rel(OUTPUT_JSON), "4h_points": data["horizons"]["4"]["common_points"], "8h_points": data["horizons"]["8"]["common_points"]}, ensure_ascii=False))

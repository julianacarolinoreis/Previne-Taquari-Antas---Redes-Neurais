#!/usr/bin/env python3
"""PREVINE AI Lab — treinamento automático leakage-safe por eventos.

Fase 1: modelos tabulares/boosting. O script nunca promove um modelo para
operação. Ele produz candidatos para modo sombra e um registro auditável.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import warnings
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from sklearn.base import clone
from sklearn.ensemble import ExtraTreesRegressor, HistGradientBoostingRegressor, RandomForestRegressor
from sklearn.linear_model import ElasticNet, Ridge
from sklearn.metrics import r2_score
from sklearn.neural_network import MLPRegressor
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.svm import SVR

try:
    from xgboost import XGBRegressor
except Exception:
    XGBRegressor = None

ROOT = Path(__file__).resolve().parents[1]
REGISTRY = ROOT / "assets" / "data" / "ai_lab" / "experiments.json"
OUT_JSON = ROOT / "assets" / "data" / "ai_lab" / "auto_training_latest.json"
OUT_CSV = ROOT / "assets" / "data" / "ai_lab" / "auto_training_leaderboard.csv"


def finite(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        x = float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None
    return x if math.isfinite(x) else None


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def parse_timestamp(values: tuple[Any, ...], header: list[str], indexes: dict[str, int], config: dict[str, Any]) -> datetime | None:
    cols = config.get("timestamp_columns")
    if cols:
        try:
            parts = [int(finite(values[indexes[name]]) or 0) for name in cols]
        except Exception:
            return None
    else:
        if len(values) < 4:
            return None
        parts = [int(finite(values[i]) or 0) for i in range(min(5, len(values)))]
    while len(parts) < 5:
        parts.append(0)
    try:
        return datetime(parts[0], parts[1], parts[2], parts[3], parts[4])
    except ValueError:
        return None


def load_rows(config: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str], dict[str, Any]]:
    source = ROOT / config["source_path"]
    if not source.exists():
        raise FileNotFoundError(source)
    wb = load_workbook(source, read_only=True, data_only=True)
    sheet = wb[config.get("sheet") or wb.sheetnames[0]]
    it = sheet.iter_rows(values_only=True)
    raw_header = list(next(it))
    header = [str(v) if v is not None else "" for v in raw_header]
    indexes = {name: i for i, name in enumerate(header) if name}
    explicit = config.get("feature_names") or []
    if explicit:
        features = [name for name in explicit if name in indexes]
    else:
        prefix = config.get("feature_prefix", "input_")
        features = [name for name in header if name.startswith(prefix)]
    if not features:
        raise RuntimeError("nenhuma feature encontrada")
    forbidden_tokens = tuple(str(x).lower() for x in config.get("forbidden_feature_tokens", ["out", "rna_final", "nivel_futuro", "target"]))
    bad = [name for name in features if any(token and token in name.lower() for token in forbidden_tokens)]
    if bad:
        raise RuntimeError("features proibidas detectadas: " + ", ".join(bad))
    required = [config["target_column"], config["event_column"], config["current_level_column"], *features]
    missing = [name for name in required if name not in indexes]
    if missing:
        raise RuntimeError("colunas ausentes: " + ", ".join(missing))

    rows: list[dict[str, Any]] = []
    skipped = 0
    for values in it:
        dt = parse_timestamp(values, header, indexes, config)
        event = finite(values[indexes[config["event_column"]]])
        target = finite(values[indexes[config["target_column"]]])
        current = finite(values[indexes[config["current_level_column"]]])
        xs = [finite(values[indexes[name]]) for name in features]
        if dt is None or event is None or target is None or current is None or any(v is None for v in xs):
            skipped += 1
            continue
        rows.append({
            "timestamp": dt,
            "event": int(event),
            "target": float(target),
            "current": float(current),
            "features": [float(v) for v in xs],
        })
    rows.sort(key=lambda r: (r["timestamp"], r["event"]))
    duplicate_keys = len(rows) - len({(r["event"], r["timestamp"]) for r in rows})
    source_hash = sha256(source)
    expected_hash = config.get("expected_source_sha256")
    if expected_hash and source_hash.lower() != str(expected_hash).lower():
        raise RuntimeError("hash da fonte divergiu do contrato do experimento")
    audit = {
        "source_path": config["source_path"],
        "source_sha256": source_hash,
        "sheet": sheet.title,
        "feature_count": len(features),
        "feature_names": features,
        "finite_rows": len(rows),
        "skipped_rows": skipped,
        "duplicate_event_timestamp_rows": duplicate_keys,
    }
    return rows, features, audit


def event_order(rows: list[dict[str, Any]]) -> list[int]:
    starts: dict[int, datetime] = {}
    for row in rows:
        starts[row["event"]] = min(starts.get(row["event"], row["timestamp"]), row["timestamp"])
    return [event for event, _ in sorted(starts.items(), key=lambda kv: kv[1])]


def build_folds(rows: list[dict[str, Any]], config: dict[str, Any]) -> list[dict[str, Any]]:
    events = event_order(rows)
    min_train_events = int(config.get("min_train_events", 5))
    min_event_rows = int(config.get("min_event_rows", 8))
    counts = {e: sum(1 for r in rows if r["event"] == e) for e in events}
    eligible = [e for e in events if counts[e] >= min_event_rows]
    folds = []
    for i in range(min_train_events, len(eligible) - 1):
        train_events = eligible[:i]
        validation_event = eligible[i]
        test_event = eligible[i + 1]
        train_rows = [r for r in rows if r["event"] in train_events]
        validation_rows = [r for r in rows if r["event"] == validation_event]
        test_rows = [r for r in rows if r["event"] == test_event]
        if not train_rows or not validation_rows or not test_rows:
            continue
        train_end = max(r["timestamp"] for r in train_rows)
        validation_start = min(r["timestamp"] for r in validation_rows)
        test_start = min(r["timestamp"] for r in test_rows)
        if not (train_end < validation_start <= test_start):
            continue
        folds.append({
            "id": f"{validation_event}_to_{test_event}",
            "train_events": train_events,
            "validation_event": validation_event,
            "test_event": test_event,
            "train_rows": train_rows,
            "validation_rows": validation_rows,
            "test_rows": test_rows,
        })
    return folds


def models(seed: int) -> dict[str, Any]:
    out: dict[str, Any] = {
        "Ridge": make_pipeline(StandardScaler(), Ridge(alpha=10.0)),
        "Elastic Net": make_pipeline(StandardScaler(), ElasticNet(alpha=0.02, l1_ratio=0.25, max_iter=5000, random_state=seed)),
        "SVR": make_pipeline(StandardScaler(), SVR(C=10.0, epsilon=0.05, gamma="scale")),
        "MLP": make_pipeline(StandardScaler(), MLPRegressor(hidden_layer_sizes=(64, 32), activation="relu", alpha=0.001, max_iter=700, early_stopping=True, validation_fraction=0.15, random_state=seed)),
        "Random Forest": RandomForestRegressor(n_estimators=260, max_depth=14, min_samples_leaf=2, n_jobs=2, random_state=seed),
        "Extra Trees": ExtraTreesRegressor(n_estimators=260, max_depth=14, min_samples_leaf=2, n_jobs=2, random_state=seed),
        "HistGradientBoosting": HistGradientBoostingRegressor(max_iter=260, learning_rate=0.05, max_leaf_nodes=31, l2_regularization=0.1, random_state=seed),
    }
    if XGBRegressor is not None:
        out["XGBoost"] = XGBRegressor(
            n_estimators=320, max_depth=4, learning_rate=0.04, min_child_weight=2,
            subsample=0.9, colsample_bytree=0.9, objective="reg:squarederror",
            eval_metric="rmse", tree_method="hist", n_jobs=2, random_state=seed,
        )
    return out


def metric(rows: list[dict[str, Any]], pred: np.ndarray) -> dict[str, Any]:
    obs = np.asarray([r["target"] for r in rows], dtype=float)
    pred = np.asarray(pred, dtype=float)
    mae = float(np.mean(np.abs(pred - obs)))
    rmse = float(np.sqrt(np.mean((pred - obs) ** 2)))
    bias = float(np.mean(pred - obs))
    sst = float(np.sum((obs - np.mean(obs)) ** 2))
    nse = float(1 - np.sum((pred - obs) ** 2) / sst) if sst else None
    r2 = float(r2_score(obs, pred)) if len(obs) > 1 else None
    oi = int(np.argmax(obs))
    pi = int(np.argmax(pred))
    peak_abs = float(abs(pred[pi] - obs[oi]))
    lag = abs((rows[pi]["timestamp"] - rows[oi]["timestamp"]).total_seconds() / 3600.0)
    return {
        "n": len(rows), "mae_cm": mae, "rmse_cm": rmse, "bias_cm": bias,
        "nse": nse, "r2": r2, "peak_abs_error_cm": peak_abs, "peak_lag_abs_h": lag,
    }


def run_experiment(config: dict[str, Any]) -> dict[str, Any]:
    rows, feature_names, audit = load_rows(config)
    folds = build_folds(rows, config)
    if len(folds) < int(config.get("min_folds", 3)):
        raise RuntimeError(f"dobras causais insuficientes: {len(folds)}")
    seed = int(config.get("seed", 20261002))
    candidates = models(seed)
    per_fold: list[dict[str, Any]] = []

    for fold in folds:
        train = fold["train_rows"]
        test = fold["test_rows"]
        x_train = np.asarray([r["features"] for r in train], dtype=float)
        y_train = np.asarray([r["target"] for r in train], dtype=float)
        x_test = np.asarray([r["features"] for r in test], dtype=float)

        persistence = metric(test, np.asarray([r["current"] for r in test], dtype=float))
        per_fold.append({"model": "Persistência", "backend": "baseline", "fold": fold["id"], **persistence})

        for name, base in candidates.items():
            estimator = clone(base)
            with warnings.catch_warnings(record=True) as caught:
                warnings.simplefilter("always")
                estimator.fit(x_train, y_train)
            pred = estimator.predict(x_test)
            m = metric(test, pred)
            per_fold.append({
                "model": name, "backend": type(estimator).__name__, "fold": fold["id"],
                "warning_count": len(caught), **m
            })

    names = sorted({r["model"] for r in per_fold})
    leaderboard = []
    persistence_mae = None
    for name in names:
        group = [r for r in per_fold if r["model"] == name]
        row = {
            "model": name,
            "fold_count": len(group),
            "median_mae_cm": float(np.median([r["mae_cm"] for r in group])),
            "mean_mae_cm": float(np.mean([r["mae_cm"] for r in group])),
            "worst_fold_mae_cm": float(np.max([r["mae_cm"] for r in group])),
            "median_rmse_cm": float(np.median([r["rmse_cm"] for r in group])),
            "median_nse": float(np.median([r["nse"] for r in group if r["nse"] is not None])),
            "median_abs_bias_cm": float(np.median([abs(r["bias_cm"]) for r in group])),
            "median_peak_abs_error_cm": float(np.median([r["peak_abs_error_cm"] for r in group])),
            "worst_peak_abs_error_cm": float(np.max([r["peak_abs_error_cm"] for r in group])),
            "median_peak_lag_abs_h": float(np.median([r["peak_lag_abs_h"] for r in group])),
            "worst_peak_lag_abs_h": float(np.max([r["peak_lag_abs_h"] for r in group])),
        }
        if name == "Persistência":
            persistence_mae = row["median_mae_cm"]
        leaderboard.append(row)

    leaderboard.sort(key=lambda r: (
        r["median_mae_cm"], r["median_rmse_cm"], r["worst_fold_mae_cm"],
        r["median_peak_abs_error_cm"], r["median_peak_lag_abs_h"]
    ))
    for rank, row in enumerate(leaderboard, 1):
        row["rank"] = rank
        row["beats_persistence"] = bool(
            persistence_mae is not None and row["model"] != "Persistência" and row["median_mae_cm"] < persistence_mae
        )
        row["shadow_eligible"] = bool(
            row["beats_persistence"]
            and row["median_nse"] > 0
            and row["fold_count"] >= int(config.get("min_folds", 3))
            and row["worst_peak_lag_abs_h"] <= float(config.get("max_peak_lag_h", max(2, int(config["horizon_hours"]) * 2)))
        )

    fold_manifest = [{
        "fold_id": f["id"],
        "train_events": f["train_events"],
        "validation_event": f["validation_event"],
        "test_event": f["test_event"],
        "train_rows": len(f["train_rows"]),
        "validation_rows": len(f["validation_rows"]),
        "test_rows": len(f["test_rows"]),
    } for f in folds]

    return {
        "schema_version": 1,
        "artifact_id": "previne_ai_lab_auto_training",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
        "experiment_id": config["id"],
        "label": config["label"],
        "station": config.get("station"),
        "target": config.get("target_name", config["target_column"]),
        "horizon_hours": config["horizon_hours"],
        "status": "completed_research_training",
        "research_only": True,
        "official_alert": False,
        "promotion_allowed": False,
        "protocol": {
            "selection": "ranking lexicográfico por MAE mediano, RMSE mediano, pior MAE, erro de pico e atraso do pico",
            "causal_event_folds": True,
            "future_columns_in_features": False,
            "baseline_required": "Persistência",
            "promotion": "nunca automática; apenas candidatos elegíveis a modo sombra",
            "phase_1_models": list(models(seed).keys()) + ["Persistência"],
            "temporal_neural_models": "fase seguinte; exigem contrato sequencial explícito",
        },
        "data_audit": audit,
        "folds": fold_manifest,
        "leaderboard": leaderboard,
        "shadow_candidates": [r for r in leaderboard if r["shadow_eligible"]],
        "shadow_queue_top3": [r for r in leaderboard if r["shadow_eligible"]][:3],
        "fold_metrics": per_fold,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--experiment", default=None)
    args = parser.parse_args()
    registry = json.loads(REGISTRY.read_text(encoding="utf-8"))
    experiments = registry["experiments"]
    config = next((x for x in experiments if x["id"] == args.experiment), None) if args.experiment else next((x for x in experiments if x.get("enabled")), None)
    if config is None:
        raise SystemExit("experimento não encontrado")
    result = run_experiment(config)
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    fields = [
        "rank","model","fold_count","median_mae_cm","mean_mae_cm","worst_fold_mae_cm",
        "median_rmse_cm","median_nse","median_abs_bias_cm","median_peak_abs_error_cm",
        "worst_peak_abs_error_cm","median_peak_lag_abs_h","worst_peak_lag_abs_h",
        "beats_persistence","shadow_eligible"
    ]
    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for row in result["leaderboard"]:
            writer.writerow({k: row.get(k) for k in fields})
    print(json.dumps({
        "experiment": result["experiment_id"],
        "models": len(result["leaderboard"]),
        "folds": len(result["folds"]),
        "shadow_candidates": [x["model"] for x in result["shadow_candidates"]],
        "output": str(OUT_JSON.relative_to(ROOT)),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()

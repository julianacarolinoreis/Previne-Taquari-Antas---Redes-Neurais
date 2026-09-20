#!/usr/bin/env python3
"""Extract source-backed Q62 event-series candidates without promoting a model.

The Q62 review package has audit workbooks but is not publishable as a selected
forecast. This builder exposes only the independent-test rows for the five
events in the declared Q62 test partition (31, 33, 34, 35 and 37), preserving
the review gate and the event-level comparison boundary.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT / "assets" / "data" / "mucum_q62_auditoria.json"
EVENTS = ROOT / "assets" / "data" / "mucum_q62" / "mucum_q62_recorte_eventos.csv"
WORKBOOKS = ROOT / "assets" / "audit_workbooks"
OUTPUT = ROOT / "assets" / "data" / "mucum_q62" / "mucum_q62_replay_candidates.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        parsed = float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None
    return parsed


def event_rows() -> dict[int, dict[str, str]]:
    import csv

    result: dict[int, dict[str, str]] = {}
    with EVENTS.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle, delimiter=";"):
            if str(row.get("usar", "")) != "1":
                continue
            try:
                event_number = int(str(row.get("evento")))
            except (TypeError, ValueError):
                continue
            result[event_number] = row
    if not result:
        raise RuntimeError("no usable events found in the Q62 event catalogue")
    return result


def declared_test_events(models: list[dict[str, Any]]) -> list[int]:
    declared: set[int] = set()
    for model in models:
        model_id = str(model.get("model_id", ""))
        match = re.search(r"_T([0-9-]+)_V", model_id)
        if not match:
            continue
        declared.update(int(value) for value in match.group(1).split("-"))
    if not declared:
        raise RuntimeError("Q62 model identifiers do not declare a test-event partition")
    return sorted(declared)


def timestamp(row: dict[str, Any]) -> str:
    return datetime(
        int(row["ANO"]), int(row["MES"]), int(row["DIA"]), int(row["HORA"])
    ).strftime("%Y-%m-%d %H:%M")


def workbook_for(model_id: str, horizon: int) -> Path:
    matches = sorted(WORKBOOKS.glob(f"{horizon}H_ALT__{model_id}.xlsx"))
    if len(matches) != 1:
        raise RuntimeError(f"expected one workbook for {model_id}, found {len(matches)}")
    return matches[0]


def event_metrics(rows: list[dict[str, Any]], horizon: int) -> dict[str, Any]:
    if not rows:
        raise RuntimeError("cannot calculate event metrics without rows")
    observed_peak = max(rows, key=lambda item: item["observed_cm"])
    predicted_peak = max(rows, key=lambda item: item["predicted_cm"])
    paired_mae = sum(abs(item["predicted_cm"] - item["observed_cm"]) for item in rows) / len(rows)
    return {
        "points": len(rows),
        "observed_peak_cm": observed_peak["observed_cm"],
        "observed_peak_base_timestamp": observed_peak["base_timestamp"],
        "observed_peak_target_timestamp": observed_peak["target_timestamp"],
        "predicted_peak_cm": predicted_peak["predicted_cm"],
        "predicted_peak_base_timestamp": predicted_peak["base_timestamp"],
        "predicted_peak_target_timestamp": predicted_peak["target_timestamp"],
        "peak_error_cm": round(predicted_peak["predicted_cm"] - observed_peak["observed_cm"], 6),
        "peak_lag_hours": round(
            (
                datetime.fromisoformat(predicted_peak["target_timestamp"])
                - datetime.fromisoformat(observed_peak["target_timestamp"])
            ).total_seconds()
            / 3600,
            6,
        ),
        "mae_cm": round(paired_mae, 6),
        "horizon_hours": horizon,
    }


def extract_model(model: dict[str, Any], review: dict[str, Any], events: dict[int, dict[str, str]]) -> dict[str, Any]:
    model_id = str(model["model_id"])
    horizon = int(str(model["horizon"]).removesuffix("h"))
    workbook = workbook_for(model_id, horizon)
    output_column = str(model["output_column"])
    required = {
        "ANO",
        "MES",
        "DIA",
        "HORA",
        "CONJUNTO",
        "COD_SEQUENCIAL",
        "EVENTO",
        "nivel_86510000",
        output_column,
        "NIVEL_FUTURO",
        "SERIE",
        "RNA_FINAL",
    }
    workbook_reader = load_workbook(workbook, read_only=True, data_only=True)
    sheet = workbook_reader[workbook_reader.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    header = list(next(iterator))
    indexes = {str(value): index for index, value in enumerate(header) if value is not None}
    missing_columns = sorted(required - set(indexes))
    if missing_columns:
        raise RuntimeError(f"{model_id}: missing columns {missing_columns}")

    rows_by_event: dict[int, list[dict[str, Any]]] = {}
    formula_mismatches = 0
    nonfinite = 0
    wrong_series = 0
    event_values: set[int] = set()
    for values in iterator:
        row = {key: values[index] for key, index in indexes.items() if index < len(values)}
        if row.get("CONJUNTO") != "Teste":
            continue
        event_number = number(row.get("EVENTO"))
        if event_number is None:
            nonfinite += 1
            continue
        event_number = int(event_number)
        event_values.add(event_number)
        if number(row.get("SERIE")) != 3:
            wrong_series += 1
            continue
        current = number(row.get("nivel_86510000"))
        target = number(row.get("NIVEL_FUTURO"))
        predicted = number(row.get("RNA_FINAL"))
        delta = number(row.get(output_column))
        if None in (current, target, predicted, delta):
            nonfinite += 1
            continue
        if abs(delta - (target - current)) > 1e-6:
            formula_mismatches += 1
        base_timestamp = timestamp(row)
        target_timestamp = (
            datetime.fromisoformat(base_timestamp) + timedelta(hours=horizon)
        ).strftime("%Y-%m-%d %H:%M")
        rows_by_event.setdefault(event_number, []).append(
            {
                "base_timestamp": base_timestamp,
                "target_timestamp": target_timestamp,
                "observed_cm": round(target, 6),
                "predicted_cm": round(predicted, 6),
            }
        )

    for event_number in rows_by_event:
        rows_by_event[event_number].sort(key=lambda item: item["base_timestamp"])
    available_events = sorted(rows_by_event)
    expected_events = sorted(events)
    if available_events != expected_events:
        raise RuntimeError(
            f"{model_id}: test events differ from Q62 catalogue; "
            f"found={available_events}, expected={expected_events}"
        )
    duplicate_timestamps_by_event = {
        str(event_number): len(rows) - len({item["base_timestamp"] for item in rows})
        for event_number, rows in rows_by_event.items()
    }
    metrics_by_event = {
        str(event_number): event_metrics(rows, horizon)
        for event_number, rows in rows_by_event.items()
    }
    default_event = 35 if 35 in rows_by_event else available_events[0]
    default_rows = rows_by_event[default_event]
    default_metrics = metrics_by_event[str(default_event)]
    review_status = str(review.get("status", "review status unavailable"))
    return {
        "model_id": model_id,
        "horizon_hours": horizon,
        "family": f"{horizon}H_ALT",
        "output_column": output_column,
        "input_names": model.get("input_names", []),
        "workbook": rel(workbook),
        "workbook_sha256": sha256(workbook),
        "review": {
            "agent": review.get("agent"),
            "status": review_status,
            "finding": review.get("finding"),
            "promotion_allowed": False,
        },
        "comparison_status": "candidate_series_extracted_not_promoted",
        "series_key": "35|Teste",
        "unit": "cm",
        "prediction_contract": {
            "observed_column": "NIVEL_FUTURO",
            "prediction_column": "RNA_FINAL",
            "target_delta_column": output_column,
            "target_and_prediction_excluded_from_inputs": all(
                forbidden not in model.get("input_names", [])
                for forbidden in ("NIVEL_FUTURO", "RNA_FINAL", output_column)
            ),
        },
        "checks": {
            "required_columns_present": not missing_columns,
            "event_values": sorted(event_values),
            "partition": "Teste",
            "serie_values": [3] if wrong_series == 0 else ["mixed_or_invalid"],
            "formula_mismatch_count": formula_mismatches,
            "nonfinite_row_count": nonfinite,
            "duplicate_timestamp_count": sum(duplicate_timestamps_by_event.values()),
            "duplicate_timestamp_count_by_event": duplicate_timestamps_by_event,
            "audit_effective_rows": model.get("effective_rows"),
            "audit_test_rows_all_events": model.get("partition_counts", {}).get("Teste"),
        },
        "events_available": available_events,
        "event_metrics": default_metrics,
        "event_metrics_by_event": metrics_by_event,
        "series": default_rows,
        "series_by_event": {str(event_number): rows for event_number, rows in rows_by_event.items()},
    }


def build() -> dict[str, Any]:
    audit = json.loads(AUDIT.read_text(encoding="utf-8"))
    catalogue = event_rows()
    models = audit.get("models", [])
    test_event_numbers = declared_test_events(models)
    missing_catalogue_events = sorted(set(test_event_numbers) - set(catalogue))
    if missing_catalogue_events:
        raise RuntimeError(f"Q62 test events missing from event catalogue: {missing_catalogue_events}")
    events = {event_number: catalogue[event_number] for event_number in test_event_numbers}
    reviews = audit.get("agentReviews", {})
    candidates: list[dict[str, Any]] = []
    for model in models:
        horizon = str(model["horizon"]).lower()
        review_models = reviews.get("h8" if horizon == "8h" else "h12", {}).get("models", {})
        review = review_models.get(model["model_id"], {})
        candidates.append(extract_model(model, review, events))
    candidates.sort(key=lambda item: (item["horizon_hours"], item["model_id"]))
    test_events = sorted({event_number for candidate in candidates for event_number in candidate["events_available"]})
    event_records = []
    for event_number in test_events:
        event = events[event_number]
        event_records.append(
            {
                "event_number": event_number,
                "peak_timestamp": event.get("pico_data"),
                "catalog_peak_cm": number(event.get("pico_cm_catalogo")),
                "observed_peak_cm": number(event.get("pico_cm_obs")),
                "recorte_start": event.get("inicio_recorte"),
                "recorte_end": event.get("fim_recorte"),
                "recorte_hours": number(event.get("horas_recorte")),
                "candidate_series_count": sum(
                    1
                    for candidate in candidates
                    if str(event_number) in candidate.get("series_by_event", {})
                ),
            }
        )
    default_event = 35 if 35 in events else test_events[0]
    default_event_row = events[default_event]
    return {
        "schema_version": 1,
        "artifact_id": "mucum-q62-replay-candidates",
        "generated_at_utc": datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "status": "research_candidates_only_not_promoted",
        "research_only": True,
        "official_alert": False,
        "municipality": "Muçum",
        "station_code": "86510000",
        "event_number": default_event,
        "event_id": f"mucum-q62-{default_event}",
        "event_peak_timestamp": default_event_row.get("pico_data"),
        "event_peak_observed_cm": number(default_event_row.get("pico_cm_obs")),
        "test_events": test_events,
        "event_count": len(test_events),
        "events": event_records,
        "recorte": {
            "start": default_event_row.get("inicio_recorte"),
            "end": default_event_row.get("fim_recorte"),
            "hours": number(default_event_row.get("horas_recorte")),
        },
        "selection": {
            "selected_model": None,
            "selection_status": "blocked_by_independent_review_gate",
            "reason": "Q62 audit reconciles MAT/XLSX/CSV and test partitions, but the independent review marked candidates conditional/not publishable; model 035 failed the performance criterion.",
        },
        "sources": {
            "audit": rel(AUDIT),
            "event_catalogue": rel(EVENTS),
        },
        "candidate_count": len(candidates),
        "candidates_by_horizon": {
            str(horizon): [item["model_id"] for item in candidates if item["horizon_hours"] == horizon]
            for horizon in (8, 12)
        },
        "candidates": candidates,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    artifact = build()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(artifact, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {args.output} · candidates={artifact['candidate_count']}")


if __name__ == "__main__":
    main()
